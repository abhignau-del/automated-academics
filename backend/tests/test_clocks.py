import io

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from automated_academics.checks import check_institution, diagnose
from automated_academics.excel_io import export_workbook, import_workbook
from automated_academics.export import build_pdf, build_xlsx
from automated_academics.grid_io import import_grid
from automated_academics.models import Calendar, LectureTime, Placement, Timetable
from automated_academics.synthetic import sample_institution
from automated_academics.views import session_views


def t(a, b):
    return LectureTime(start=a, end=b)


LONG = [t(f"{8 + i}:00", f"{8 + i}:50") for i in range(7)]  # 50-minute lectures on the hour
SHORT = [t(f"{8 + i}:00", f"{8 + i}:40") for i in range(7)]  # 40-minute ones


def with_clocks():
    inst = sample_institution()
    inst.calendar = inst.calendar.model_copy(update={"times": LONG, "clocks": {"40-minute": SHORT}})
    inst.batches = [b.model_copy(update={"clock": "40-minute"}) if b.id.startswith("MG-") else b for b in inst.batches]
    return inst


# ---------- the model ----------

def test_a_class_resolves_its_clock_then_the_day_then_the_default():
    cal = Calendar(day_names=["Mon", "Tue"], lectures_per_day=2, break_after=[], times=[t("9:00", "10:00"), t("10:00", "11:00")],
                   clocks={"short": [t("9:00", "9:40"), t("9:45", "10:25")]}, day_times={1: [t("13:00", "13:30"), t("13:30", "14:00")]})
    assert cal.time_label(0, 0) == "9:00–10:00"
    assert cal.time_label(0, 0, 1, "short") == "9:00–9:40"
    assert cal.time_label(1, 0, 1, "short") == "13:00–13:30"  # a day that runs differently for everyone wins
    assert cal.time_label(0, 0, 1, "no such clock") == "9:00–10:00"  # an unknown name falls back to the default


def test_a_clock_needs_a_time_for_every_lecture_in_order():
    with pytest.raises(ValidationError, match="the clock 'x'"):
        Calendar(lectures_per_day=3, break_after=[], clocks={"x": [t("9:00", "10:00")]})
    with pytest.raises(ValidationError, match="starts before lecture 1 ends"):
        Calendar(lectures_per_day=2, break_after=[], clocks={"x": [t("9:00", "10:00"), t("9:30", "10:30")]})
    with pytest.raises(ValidationError, match="needs a name"):
        Calendar(lectures_per_day=1, break_after=[], clocks={" ": [t("9:00", "10:00")]})


def test_a_class_must_name_a_clock_that_exists():
    inst = sample_institution()
    inst.batches[0] = inst.batches[0].model_copy(update={"clock": "ghost"})
    data = inst.model_dump(mode="json")
    parsed, issues = check_institution(data)
    assert parsed is None
    assert any(i.section == "batches" and i.index == 0 and i.field == "clock" for i in issues)


def test_a_practical_group_keeps_its_parent_clock():
    inst = with_clocks()
    inst.batches = [b.model_copy(update={"clock": "40-minute"}) if b.id == "CS-UG1" else b for b in inst.batches]
    assert inst.clock_of_batch("CS-UG1-A") == "40-minute"  # the group has none of its own
    assert inst.clock_of_batch("CS-UG3-A") is None
    assert inst.clock_of_batch("nobody") is None


# ---------- what people see ----------

def test_sessions_carry_their_own_classes_times():
    inst = with_clocks()
    off_mg = next(o for o in inst.offerings if o.batch_ids == ["MG-UG1"])
    off_cs = next(o for o in inst.offerings if o.batch_ids == ["CS-UG3"])
    tt = Timetable(status="x", placements=[
        Placement(offering_id=off_mg.id, session_index=0, day=0, start=1, length=1, room_id="C1"),
        Placement(offering_id=off_cs.id, session_index=0, day=0, start=1, length=1, room_id="C2"),
    ])
    mg = session_views(inst, tt, "batch", "MG-UG1")
    cs = session_views(inst, tt, "batch", "CS-UG3")
    assert mg[0]["time"] == "9:00–9:40" and cs[0]["time"] == "9:00–9:50"
    # a teacher's or room's view carries each session's own time too
    room = session_views(inst, tt, "room", "C1")
    assert room[0]["time"] == "9:00–9:40"


def grid_cells(ws):
    return {(r, c): ws.cell(row=r, column=c).value for r in range(1, 12) for c in range(1, 9)}


def test_a_classs_excel_grid_is_labelled_with_its_own_clock_and_others_show_theirs():
    inst = with_clocks()
    off_cs = next(o for o in inst.offerings if o.batch_ids == ["CS-UG3"])
    off_mg = next(o for o in inst.offerings if o.batch_ids == ["MG-UG1"])
    tt = Timetable(status="x", placements=[
        Placement(offering_id=off_cs.id, session_index=0, day=0, start=0, length=1, room_id="C1"),
        Placement(offering_id=off_mg.id, session_index=0, day=1, start=1, length=1, room_id="C2"),
    ])
    buf = io.BytesIO()
    build_xlsx(inst, tt, buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    by_title = {ws.title: ws for ws in wb.worksheets}
    mg = next(ws for n, ws in by_title.items() if "MG-UG1" in n)
    cs = next(ws for n, ws in by_title.items() if "CS-UG3" in n and "CS-UG3-" not in n)
    assert mg.cell(row=3, column=1).value == "L1\n8:00–8:40"  # MG-UG1 keeps the 40-minute clock
    assert cs.cell(row=3, column=1).value == "L1\n8:00–8:50"
    teacher = next(ws for n, ws in by_title.items() if n.startswith("Faculty"))
    # a teacher's grid is labelled with the default; a session on another clock says so
    mg_teacher = next(f.id for f in inst.faculty if f.id == off_mg.faculty_id)
    tws = next(ws for n, ws in by_title.items() if mg_teacher in n)
    cells = [v for v in grid_cells(tws).values() if isinstance(v, str) and "9:00–9:40" in v]
    assert cells, "the 40-minute session should show its own time on a teacher's grid"


def test_a_pdf_page_shows_the_classs_clock(tmp_path):
    from pypdf import PdfReader
    inst = with_clocks()
    tt = Timetable(status="x", placements=[])
    buf = io.BytesIO()
    off_mg = next(o for o in inst.offerings if o.batch_ids == ["MG-UG1"])
    tt = Timetable(status="x", placements=[Placement(offering_id=off_mg.id, session_index=0, day=0, start=0, length=1, room_id="C1")])
    build_pdf(inst, tt, buf, "batch", "MG-UG1")
    assert "8:40" in PdfReader(io.BytesIO(buf.getvalue())).pages[0].extract_text()


# ---------- Excel ----------

def test_clocks_and_class_clocks_survive_the_excel_round_trip(tmp_path):
    inst = with_clocks()
    path = tmp_path / "c.xlsx"
    export_workbook(inst, path)
    back = import_workbook(path)
    assert back.calendar.clocks == {"40-minute": SHORT}
    assert {b.id: b.clock for b in back.batches if b.clock} == {b.id: "40-minute" for b in inst.batches if b.clock}
    cells = {r[0].value: r[1].value for r in load_workbook(path)["Institution"].iter_rows(min_row=2)}
    assert cells["clock_40-minute"].startswith("8:00-8:40; 9:00-9:40")


def test_an_old_workbook_without_clocks_imports(tmp_path):
    path = tmp_path / "old.xlsx"
    export_workbook(sample_institution(), path)
    wb = load_workbook(path)
    ws = wb["Batches"]
    assert ws.cell(row=1, column=9).value == "clock"
    ws.delete_cols(9)  # a workbook written before clocks existed
    wb.save(path)
    assert import_workbook(path).calendar.clocks == {}


# ---------- the warning ----------

def test_classes_on_different_clocks_that_share_a_teacher_are_warned_about():
    inst = with_clocks()
    # a clock offset by half a lecture: lecture 2 of one runs into lecture 1 of the other
    shifted = [t(f"{8 + i}:30", f"{9 + i}:10") for i in range(7)]
    inst.calendar = inst.calendar.model_copy(update={"clocks": {"40-minute": shifted}})
    shared = inst.faculty[0].id
    mg = next(o for o in inst.offerings if o.batch_ids == ["MG-UG1"])
    inst.offerings = [o.model_copy(update={"faculty_id": shared}) if o.id == mg.id else o for o in inst.offerings]
    assert any("share teachers" in d.message and d.level == "warning" for d in diagnose(inst))


def test_aligned_clocks_or_no_shared_teacher_give_no_warning():
    inst = with_clocks()  # same start times, only the ends differ; each department has its own teachers
    assert not any("share teachers" in d.message for d in diagnose(inst))


# ---------- importing grids on different clocks ----------

DAYS = ["MON", "TUE", "WED"]


def put(ws, top, title, rows, times):
    ws.cell(row=top, column=1, value=title)
    for j, d in enumerate(DAYS):
        ws.cell(row=top + 1, column=2 + j, value=d)
    for i, row in enumerate(rows):
        ws.cell(row=top + 2 + i, column=1, value=times[i])
        for j, v in enumerate(row):
            if v:
                ws.cell(row=top + 2 + i, column=2 + j, value=v)


def two_clock_workbook():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("S")
    fifty = ["8 TO 8.50", "8.55 TO 9.45"]
    forty = ["8 TO 8.40", "8.45 TO 9.25"]
    put(ws, 1, "Class A", [["M(Prof. A One)", "N(Prof. B Two)", ""], ["N(Prof. B Two)", "", "M(Prof. A One)"]], fifty)
    put(ws, 6, "Class B", [["P(Prof. C Three)", "", "Q(Prof. D Four)"], ["", "P(Prof. C Three)", ""]], fifty)
    put(ws, 11, "Class C", [["R(Prof. E Five)", "", ""], ["", "S(Prof. F Six)", ""]], forty)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_classes_on_a_different_clock_get_their_own_named_clock():
    inst, report, _ = import_grid(two_clock_workbook())
    assert [x.end for x in inst.calendar.times] == ["08:50", "09:45"]  # the most common set is the default
    assert list(inst.calendar.clocks) == ["40-minute lectures"]
    clocks = {b.program: b.clock for b in inst.batches}
    assert clocks == {"Class A": None, "Class B": None, "Class C": "40-minute lectures"}
    assert any("keep other hours" in a for a in report.assumptions)
    assert not any("can't differ by class" in p for p in report.problems)


def test_the_model_itself_refuses_an_unknown_clock():
    inst = sample_institution()
    data = inst.model_dump(mode="json")
    data["batches"][0]["clock"] = "ghost"
    with pytest.raises(ValidationError, match="unknown clock"):
        type(inst).model_validate(data)


def test_a_class_page_in_the_pdf_is_labelled_with_its_own_clock_not_the_default():
    from pypdf import PdfReader
    inst = with_clocks()  # default lectures end at :50, MG- classes' at :40
    buf = io.BytesIO()
    build_pdf(inst, Timetable(status="x", placements=[]), buf, "batch", "MG-UG1")
    mg = PdfReader(io.BytesIO(buf.getvalue())).pages[0].extract_text()
    assert "8:40" in mg and "8:50" not in mg
    buf = io.BytesIO()
    build_pdf(inst, Timetable(status="x", placements=[]), buf, "batch", "CS-UG3")
    cs = PdfReader(io.BytesIO(buf.getvalue())).pages[0].extract_text()
    assert "8:50" in cs and "8:40" not in cs
