import io

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from pypdf import PdfReader

from automated_academics.checks import check_institution
from automated_academics.excel_io import export_workbook, import_workbook
from automated_academics.export import build_pdf, build_xlsx
from automated_academics.models import Calendar, LectureTime
from automated_academics.solver import solve
from automated_academics.synthetic import sample_institution
from automated_academics.views import session_views


def t(a, b):
    return LectureTime(start=a, end=b)


FOUR = [t("8:00", "8:50"), t("8:55", "9:45"), t("9:50", "10:40"), t("11:00", "11:50")]


def cal(**kw):
    return Calendar(day_names=["Mon", "Tue", "Wed"], lectures_per_day=4, break_after=[2], **kw)


# ---------- the model ----------

def test_times_are_normalised_and_prettified():
    x = t("9:30", "10:30")
    assert (x.start, x.end) == ("09:30", "10:30")
    assert LectureTime.pretty("09:30") == "9:30" and LectureTime.pretty("14:05") == "14:05"
    assert cal(times=FOUR).time_label(0, 0) == "8:00–8:50"
    assert cal(times=FOUR).time_label(1, 1, 2) == "8:55–10:40"  # a two-lecture block spans both


def test_no_times_means_no_labels():
    c = cal()
    assert c.times == [] and c.time_label(0, 0) is None
    assert Calendar().times == []  # old data, with no times, still loads


@pytest.mark.parametrize("bad", ["", "9", "25:00", "9:60", "nine thirty", "9.30pm"])
def test_nonsense_clock_times_are_rejected(bad):
    with pytest.raises(ValidationError):
        LectureTime(start=bad, end="23:00")


def test_a_lecture_must_end_after_it_starts():
    with pytest.raises(ValidationError):
        t("10:00", "10:00")
    with pytest.raises(ValidationError):
        t("10:30", "09:30")


def test_one_time_for_every_lecture_in_order():
    with pytest.raises(ValidationError, match="each of the 4 lectures"):
        cal(times=FOUR[:3])
    overlapping = [t("8:00", "9:00"), t("8:30", "9:30"), t("9:30", "10:00"), t("10:00", "10:30")]
    with pytest.raises(ValidationError, match="starts before lecture 1 ends"):
        cal(times=overlapping)


def test_a_day_can_run_to_its_own_clock():
    short = [t("8:00", "8:40"), t("8:45", "9:25"), t("9:35", "10:15"), t("10:20", "11:00")]
    c = cal(times=FOUR, day_times={2: short})
    assert c.time_label(0, 3) == "11:00–11:50"
    assert c.time_label(2, 3) == "10:20–11:00"
    with pytest.raises(ValidationError, match="not one of the 3 days"):
        cal(times=FOUR, day_times={5: short})
    with pytest.raises(ValidationError, match="Wed"):
        cal(times=FOUR, day_times={2: short[:2]})


def test_json_round_trip_keeps_day_overrides():
    c = cal(times=FOUR, day_times={1: FOUR})
    again = Calendar.model_validate_json(c.model_dump_json())
    assert again == c and list(again.day_times) == [1]


# ---------- located errors in the editor ----------

def test_a_bad_time_is_reported_against_settings():
    data = sample_institution().model_dump(mode="json")
    data["calendar"]["times"] = [{"start": "9:30", "end": "10:30"}]  # one row for 7 lectures
    inst, issues = check_institution(data)
    assert inst is None
    assert any(i.section == "settings" and "each of the 7 lectures" in i.message for i in issues)


# ---------- Excel ----------

def test_times_survive_the_excel_round_trip(tmp_path):
    inst = sample_institution()
    seven = [t(f"{8 + i}:00", f"{8 + i}:50") for i in range(7)]
    inst.calendar = inst.calendar.model_copy(update={"times": seven, "day_times": {5: [t(f"{8 + i}:00", f"{8 + i}:30") for i in range(7)]}})
    path = tmp_path / "t.xlsx"
    export_workbook(inst, path)
    back = import_workbook(path)
    assert back.calendar.times == seven
    assert back.calendar.day_times == inst.calendar.day_times
    ws = load_workbook(path)["Institution"]
    cells = {r[0].value: r[1].value for r in ws.iter_rows(min_row=2)}
    assert cells["lecture_times"].startswith("8:00-8:50; 9:00-9:50")
    assert "lecture_times_Sat" in cells


def test_a_workbook_without_times_still_imports(tmp_path):
    path = tmp_path / "old.xlsx"
    export_workbook(sample_institution(), path)
    assert import_workbook(path).calendar.times == []


def test_people_can_type_times_in_their_own_style(tmp_path):
    path = tmp_path / "t.xlsx"
    export_workbook(sample_institution(), path)
    wb = load_workbook(path)
    ws = wb["Institution"]
    ws.append(["lecture_times", "8 to 8.50; 9:00 - 9:50 ; 10-10.50; 11:00-11:50; 12:00 - 12:50; 1:00 pm - 1:50 pm; 2.00pm-2.50pm"])
    wb.save(path)
    back = import_workbook(path)
    assert [x.start for x in back.calendar.times] == ["08:00", "09:00", "10:00", "11:00", "12:00", "13:00", "14:00"]


# ---------- what people see ----------

@pytest.fixture(scope="module")
def timed():
    inst = sample_institution()
    seven = [t("8:00", "8:50"), t("8:55", "9:45"), t("9:50", "10:40"), t("10:45", "11:35"),
             t("12:30", "13:20"), t("13:25", "14:15"), t("14:20", "15:10")]
    inst.calendar = inst.calendar.model_copy(update={"times": seven, "day_times": {5: [t(f"{8 + i}:00", f"{8 + i}:40") for i in range(7)]}})
    return inst, solve(inst, time_limit_s=4)


def test_session_views_carry_the_time(timed):
    inst, tt = timed
    sessions = session_views(inst, tt, "batch", "CS-UG1")
    assert sessions and all(s["time"] == inst.calendar.time_label(s["day"], s["start"], s["length"]) for s in sessions)
    assert all(s["time"] and "–" in s["time"] for s in sessions)


def test_exports_show_times(timed):
    inst, tt = timed
    buf = io.BytesIO()
    build_xlsx(inst, tt, buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    grid = next(ws for ws in wb.worksheets if ws.title != wb.worksheets[0].title)
    labels = [grid.cell(row=r, column=1).value for r in range(3, 10)]
    assert labels[0] == "L1\n8:00–8:50" and labels[4] == "L5\n12:30–13:20"

    pdf = io.BytesIO()
    build_pdf(inst, tt, pdf, "batch", "CS-UG1")
    text = PdfReader(io.BytesIO(pdf.getvalue())).pages[0].extract_text()
    assert "8:00" in text and "12:30" in text


def test_without_times_exports_look_as_before():
    inst = sample_institution()
    tt = solve(inst, time_limit_s=4)
    buf = io.BytesIO()
    build_xlsx(inst, tt, buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    grid = wb.worksheets[1]
    assert grid.cell(row=3, column=1).value == "L1"
    assert all(s["time"] is None for s in session_views(inst, tt, "batch", "CS-UG1"))


@pytest.mark.parametrize("typed, start, end", [
    ("8 TO 8.50", "08:00", "08:50"),
    ("8.55 to 9.45", "08:55", "09:45"),
    ("9:30 am - 10:30 am", "09:30", "10:30"),
    ("1:15 pm \u2013 2:05 pm", "13:15", "14:05"),
    ("11-11.50", "11:00", "11:50"),
])
def test_the_ways_timetables_are_written(typed, start, end):
    from automated_academics.excel_io import _times
    (got,) = _times(typed, "x")
    assert (got.start, got.end) == (start, end)


def test_a_block_that_runs_off_the_day_has_no_label():
    c = cal(times=FOUR)
    assert c.time_label(0, 3, 1) == "11:00\u201311:50"
    assert c.time_label(0, 3, 2) is None  # would need a fifth lecture
    assert c.time_label(0, -1) is None


def test_a_day_with_its_own_clock_shows_it_in_the_excel_cell(timed):
    from automated_academics.models import Placement, Timetable

    inst, _ = timed
    # placed by hand (not solved) so there is certainly a session on Saturday, the day with its own clock, and one on Monday
    off = next(o for o in inst.offerings if o.batch_ids == ["CS-UG1"])
    tt = Timetable(status="MANUAL", placements=[
        Placement(offering_id=off.id, session_index=0, day=5, start=1, length=1, room_id="C1"),
        Placement(offering_id=off.id, session_index=1, day=0, start=1, length=1, room_id="C1"),
    ])
    buf = io.BytesIO()
    build_xlsx(inst, tt, buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    grid = next(ws for ws in wb.worksheets if "CS-UG1" in ws.title and "UG1-" not in ws.title)
    sat = grid.cell(row=4, column=7).value  # Saturday is the sixth day (column G); lecture 2 is row 4
    mon = grid.cell(row=4, column=2).value
    assert sat and sat.split("\n")[0] == "9:00\u20139:40"  # Saturday's own clock, from day_times
    assert mon and "\u2013" not in mon.split("\n")[0]  # Monday keeps the left-hand labels only
