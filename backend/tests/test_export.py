import io
import os
import time

import pytest
import reportlab
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from automated_academics.api import _attachment, create_app
from automated_academics.export import _selection, _sheet_name, build_pdf, build_xlsx, layout
from automated_academics.solver import solve
from automated_academics.synthetic import sample_institution
from automated_academics.views import session_views


@pytest.fixture(scope="module")
def solved():
    inst = sample_institution()
    return inst, solve(inst, time_limit_s=4)


def sess(day, start, length, code="C1", oid="O1"):
    return {"day": day, "start": start, "length": length, "course_code": code, "offering_id": oid}


# ---------- layout ----------

def test_layout_merges_exclusive_multi_lecture_sessions_only():
    lab, parallel_a, parallel_b = sess(0, 4, 2, oid="A"), sess(1, 4, 2, oid="B"), sess(1, 4, 2, oid="C")
    cover, merge = layout([lab, parallel_a, parallel_b])
    assert merge == [0]  # parallel groups share cells, so they cannot be merged
    assert cover[(1, 4)] == [1, 2] and cover[(0, 5)] == [0]


def test_layout_partial_overlap_is_not_merged():
    cover, merge = layout([sess(0, 0, 2, oid="A"), sess(0, 1, 1, oid="B")])
    assert merge == [] and cover[(0, 1)] == [0, 1]


# ---------- Excel ----------

def test_sheet_names_are_valid_and_unique():
    taken: set[str] = set()
    a = _sheet_name("Class", "CS/UG:1?", taken)
    b = _sheet_name("Class", "CS/UG:1?", taken)
    long1 = _sheet_name("Faculty", "x" * 60, taken)
    long2 = _sheet_name("Faculty", "x" * 60, taken)
    assert a == "Class CS-UG-1-" and b != a
    assert len(long1) <= 31 and len(long2) <= 31 and long1 != long2
    assert not any(ch in a + b + long1 for ch in "[]:*?/\\")


def test_xlsx_structure_and_content(solved):
    inst, tt = solved
    buf = io.BytesIO()
    build_xlsx(inst, tt, buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()))

    n_batches, n_fac, n_rooms = len(inst.batches), len(inst.faculty), len(inst.rooms)
    assert wb.sheetnames[0] == "All sessions"
    assert {"Class CS-UG1", "Faculty CSF1", "Room L1"} <= set(wb.sheetnames)
    assert len(wb.sheetnames) - 1 <= n_batches + n_fac + n_rooms

    flat = wb["All sessions"]
    assert flat.max_row - 1 == len(tt.placements)  # every session exactly once

    ws = wb["Class CS-UG1"]
    sessions = session_views(inst, tt, "batch", "CS-UG1")
    for s in sessions:
        cell = ws.cell(row=s["start"] + 3, column=s["day"] + 2)
        assert str(cell.value).startswith(s["course_code"]), (s, cell.value)
    assert ws.merged_cells.ranges  # the multi-lecture labs are merged


def test_xlsx_parallel_groups_listed_in_one_cell():
    inst = sample_institution()
    from automated_academics.models import Placement, Timetable
    # force both CS-UG1 lab groups into the same slot (it is a display test, not a valid timetable)
    labs = [o for o in inst.offerings if o.id.startswith("O-CS-UG1-LAB-")]
    tt = Timetable(status="MANUAL", placements=[
        Placement(offering_id=o.id, session_index=0, day=0, start=4, length=2, room_id=r)
        for o, r in zip(labs, ("L1", "L2"))])
    buf = io.BytesIO()
    build_xlsx(inst, tt, buf)
    ws = load_workbook(io.BytesIO(buf.getvalue()))["Class CS-UG1"]
    text = ws.cell(row=4 + 3, column=2).value
    assert text.count("\n") == 1 and text.count("CS-UG1-LAB") == 2


# ---------- PDF ----------

def pdf_pages(inst, tt, **kw):
    buf = io.BytesIO()
    build_pdf(inst, tt, buf, **kw)
    assert buf.getvalue().startswith(b"%PDF")
    return PdfReader(io.BytesIO(buf.getvalue())).pages


def test_pdf_has_one_page_per_nonempty_entity(solved):
    inst, tt = solved
    pages = pdf_pages(inst, tt)
    assert len(pages) == len(_selection(inst, tt, None, None))
    assert len(pages) >= len(inst.batches) + len(inst.faculty)
    first = pages[0].extract_text()
    assert "Class: CS-UG1" in first and inst.name in first and "Mon" in first


def test_pdf_single_entity_and_filters(solved):
    inst, tt = solved
    pages = pdf_pages(inst, tt, kind="faculty", ident="CSF1")
    assert len(pages) == 1 and "CS Faculty 1" in pages[0].extract_text()
    assert len(pdf_pages(inst, tt, kind="batch")) == len(inst.batches)
    with pytest.raises(LookupError):
        pdf_pages(inst, tt, kind="room", ident="NOPE")


def test_pdf_requested_empty_entity_still_renders(solved):
    inst, tt = solved
    empty = tt.model_copy(update={"placements": []})
    pages = pdf_pages(inst, empty, kind="room", ident="H1")
    assert len(pages) == 1
    assert len(pdf_pages(inst, empty)) == 1  # "No sessions to export."


VERA = os.path.join(os.path.dirname(reportlab.__file__), "fonts", "Vera.ttf")  # ships with reportlab


def test_pdf_uses_configured_font(solved):
    inst, tt = solved
    plain, custom = io.BytesIO(), io.BytesIO()
    build_pdf(inst, tt, plain, "batch", "CS-UG1")
    build_pdf(inst, tt, custom, "batch", "CS-UG1", font_path=VERA)
    assert b"Vera" not in plain.getvalue() and b"Vera" in custom.getvalue()
    assert "Class: CS-UG1" in PdfReader(io.BytesIO(custom.getvalue())).pages[0].extract_text()


def test_pdf_warns_when_text_needs_a_font_that_is_not_configured(solved, caplog):
    inst, tt = solved
    inst = inst.model_copy(deep=True)
    inst.faculty[0].name = "डॉ. शर्मा"
    with caplog.at_level("WARNING", logger="automated_academics"):
        build_pdf(inst, tt, io.BytesIO(), "batch", "CS-UG1")
    assert "AA_PDF_FONT" in caplog.text
    caplog.clear()
    with caplog.at_level("WARNING", logger="automated_academics"):
        build_pdf(inst, tt, io.BytesIO(), "batch", "CS-UG1", font_path=VERA)  # configured: no warning
        build_pdf(*solved, io.BytesIO(), "batch", "CS-UG1")  # all-Latin data: no warning
    assert caplog.text == ""


def test_bad_font_path_fails_at_startup(tmp_path):
    with pytest.raises(FileNotFoundError):
        create_app(str(tmp_path / "f.db"), pdf_font=str(tmp_path / "missing.ttf"))


def test_pdf_escapes_markup_in_names(solved):
    inst, tt = solved
    inst = inst.model_copy(deep=True)
    inst.courses[0].name = "Maths <b>& Stats</i>"
    pages = pdf_pages(inst, tt, kind="batch", ident="CS-UG1")
    assert "Maths <b>& Stats</i>" in pages[0].extract_text()


# ---------- API ----------

@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(str(tmp_path / "e.db"))) as c:
        yield c


def solved_job(client, inst=None):
    inst = inst or sample_institution()
    iid = client.post("/institutions", json=inst.model_dump(mode="json")).json()["id"]
    jid = client.post(f"/institutions/{iid}/solve", json={"time_limit_s": 4}).json()["job_id"]
    end = time.time() + 90
    while time.time() < end:
        job = client.get(f"/jobs/{jid}").json()
        if job["status"] in ("done", "failed"):
            return iid, jid, job["status"]
        time.sleep(0.5)
    raise AssertionError("job did not finish")


def test_export_endpoints(client):
    _, jid, status = solved_job(client)
    assert status == "done"

    x = client.get(f"/jobs/{jid}/export.xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    assert 'filename="timetable.xlsx"' in x.headers["content-disposition"]
    assert "All sessions" in load_workbook(io.BytesIO(x.content)).sheetnames

    p = client.get(f"/jobs/{jid}/export.pdf")
    assert p.status_code == 200 and p.headers["content-type"] == "application/pdf" and p.content.startswith(b"%PDF")

    one = client.get(f"/jobs/{jid}/export.pdf", params={"kind": "batch", "id": "MG-UG1"})
    assert len(PdfReader(io.BytesIO(one.content)).pages) == 1
    assert 'filename="timetable-batch-MG-UG1.pdf"' in one.headers["content-disposition"]

    assert client.get(f"/jobs/{jid}/export.pdf", params={"kind": "batch", "id": "NOPE"}).status_code == 404
    assert client.get(f"/jobs/{jid}/export.pdf", params={"id": "MG-UG1"}).status_code == 422
    assert client.get(f"/jobs/{jid}/export.pdf", params={"kind": "bogus"}).status_code == 422
    assert client.get("/jobs/nope/export.xlsx").status_code == 404


def test_export_unfinished_job_is_409(client):
    inst = sample_institution()
    inst.rooms = [r for r in inst.rooms if r.kind.value != "lab"]
    _, jid, status = solved_job(client, inst)
    assert status == "failed"
    assert client.get(f"/jobs/{jid}/export.xlsx").status_code == 409
    assert client.get(f"/jobs/{jid}/export.pdf").status_code == 409


def test_export_survives_a_saved_edit_with_bad_references(client):
    _, jid, _ = solved_job(client)
    tt = client.get(f"/jobs/{jid}/timetable").json()
    tt["placements"][0]["room_id"] = "GONE"
    tt["placements"][1]["day"] = 99
    assert client.put(f"/jobs/{jid}/timetable", json=tt).status_code == 200
    assert client.get(f"/jobs/{jid}/export.xlsx").status_code == 200
    assert client.get(f"/jobs/{jid}/export.pdf").status_code == 200


def test_attachment_filename_is_sanitised():
    header = _attachment('time"table\r\nX-Evil: 1/../x.pdf')["Content-Disposition"]
    assert "\r" not in header and "\n" not in header and header.count('"') == 2
