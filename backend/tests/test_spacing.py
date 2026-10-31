"""Course spacing: pairs of consecutive days on which the same course meets."""

import io
import logging

from automated_academics.excel_io import export_workbook, import_workbook
from automated_academics.models import Pin, Weights
from automated_academics.quality import METRICS, measure, score
from automated_academics.solver import solve
from automated_academics.synthetic import sample_institution
from test_quality import ZERO, place, tiny, tt

SPACING = Weights(**{**ZERO.model_dump(), "day_spacing": 4})


# ---------- the measurement ----------

def test_it_counts_pairs_of_consecutive_days_per_course():
    inst = tiny([("O1", "F1", ["B"], [1, 1, 1]), ("O2", "F2", ["B"], [1, 1])], days=5)
    def spacing(*ps):
        return measure(inst, tt(*ps))["day_spacing"]
    assert spacing(place("O1", 0, 0), place("O1", 2, 0, session=1), place("O1", 4, 0, session=2)) == 0   # Mon, Wed, Fri
    assert spacing(place("O1", 0, 0), place("O1", 1, 0, session=1), place("O1", 4, 0, session=2)) == 1   # Mon+Tue
    assert spacing(place("O1", 0, 0), place("O1", 1, 0, session=1), place("O1", 2, 0, session=2)) == 2   # Mon, Tue, Wed
    assert spacing(place("O1", 0, 0), place("O1", 4, 0, session=1)) == 0                                  # the week does not wrap


def test_two_sessions_on_one_day_count_as_one_day_and_courses_are_counted_separately():
    inst = tiny([("O1", "F1", ["B"], [1, 1, 1]), ("O2", "F2", ["B"], [1, 1])], days=5)
    ps = [place("O1", 0, 0), place("O1", 0, 2, session=1), place("O1", 1, 0, session=2),  # O1: Mon (twice), Tue
          place("O2", 1, 4, room="R1"), place("O2", 3, 4, session=1, room="R1")]            # O2: Tue, Thu
    assert measure(inst, tt(*ps))["day_spacing"] == 1  # only O1's Mon+Tue; O1 and O2 meeting on adjacent days is nothing
    assert measure(inst, tt(*ps[:3], place("O2", 2, 4, room="R1"), place("O2", 3, 4, session=1, room="R1")))["day_spacing"] == 2


def test_it_is_a_named_metric_with_a_weight_and_is_off_by_default():
    assert "day_spacing" in METRICS and Weights().day_spacing == 0
    assert score(Weights(day_spacing=3), {"day_spacing": 2}) == 6
    assert score(Weights(), {"day_spacing": 9}) == 0


# ---------- the solver ----------

def test_the_solver_spreads_a_courses_sessions_over_days_when_asked():
    inst = tiny([("O1", "F1", ["B"], [1, 1])], days=5, weights=SPACING)
    together = [place("O1", 0, 1), place("O1", 1, 1, session=1)]
    assert measure(inst, tt(*together))["day_spacing"] == 1
    out = solve(inst, time_limit_s=10, workers=1, hint=together, min_first_solution_s=5)
    assert out.breakdown["day_spacing"] == 0
    days = sorted(p.day for p in out.placements)
    assert days[1] - days[0] != 1


def test_it_does_nothing_when_the_goal_is_off():
    inst = tiny([("O1", "F1", ["B"], [1, 1])], days=5, weights=ZERO)
    together = [place("O1", 0, 1), place("O1", 1, 1, session=1)]
    out = solve(inst, time_limit_s=10, workers=1, hint=together, min_first_solution_s=5)
    assert out.status in ("OPTIMAL", "FEASIBLE")  # every goal off: the first timetable is kept as it is
    assert out.breakdown["day_spacing"] == measure(inst, out)["day_spacing"]


def test_it_accepts_what_cannot_be_avoided_and_never_does_worse_than_forced():
    inst = tiny([("O1", "F1", ["B"], [1, 1, 1])], days=3, weights=SPACING)
    inst = inst.model_copy(update={"pins": [Pin(offering_id="O1", session_index=0, day=0, start=0),
                                            Pin(offering_id="O1", session_index=1, day=1, start=0)]})
    out = solve(inst, time_limit_s=10, workers=1)
    # sessions 1 and 2 are pinned to Mon and Tue (one pair); the third can sit on Mon or Tue as well, adding no new pair
    assert out.breakdown["day_spacing"] == 1
    assert {p.day for p in out.placements} <= {0, 1}


def test_the_model_and_the_measurement_agree_on_the_sample(caplog):
    base = sample_institution()
    inst = base.model_copy(update={"weights": base.weights.model_copy(update={"day_spacing": 3})})
    with caplog.at_level(logging.WARNING, logger="automated_academics"):
        out = solve(inst, time_limit_s=20, workers=2)
    assert not [r for r in caplog.records if "model says" in r.getMessage()]  # the solver's own self-check found no mismatch
    assert out.breakdown["day_spacing"] == measure(inst, out)["day_spacing"]
    plain = solve(base, time_limit_s=20, workers=2)
    assert out.breakdown["day_spacing"] <= plain.breakdown["day_spacing"]


# ---------- the workbook ----------

def test_the_weight_survives_the_workbook_and_a_missing_row_means_off():
    inst = sample_institution()
    inst.weights.day_spacing = 7
    buf = io.BytesIO()
    export_workbook(inst, buf)
    assert import_workbook(io.BytesIO(buf.getvalue())).weights.day_spacing == 7
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
    ws = wb["Institution"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "weight_day_spacing")
    ws.delete_rows(row)
    out = io.BytesIO()
    wb.save(out)
    assert import_workbook(io.BytesIO(out.getvalue())).weights.day_spacing == 0


# ---------- the checks ----------

def test_checks_say_when_courses_meet_too_often_to_be_kept_apart():
    from automated_academics.checks import diagnose
    def notes(inst):
        return [d for d in diagnose(inst) if "course-spacing" in d.message]
    # a 5-day week keeps at most 3 days apart: a 4-session course must double up or meet on consecutive days once
    inst = tiny([("O1", "F1", ["B"], [1, 1, 1, 1]), ("O2", "F2", ["B"], [1, 1, 1]), ("O3", "F3", ["B"], [1, 1, 1, 1, 1])], days=5, weights=SPACING)
    found = notes(inst)
    assert len(found) == 1 and found[0].level == "info"
    assert "2 course(s)" in found[0].message and "at least 3 of their sessions" in found[0].message and "most for O3: 2" in found[0].message
    assert not notes(inst.model_copy(update={"weights": ZERO}))  # goal off: nothing said
    calm = tiny([("O1", "F1", ["B"], [1, 1, 1])], days=5, weights=SPACING)
    assert not notes(calm)  # three sessions fit on Mon, Wed, Fri
    even = tiny([("O1", "F1", ["B"], [1, 1, 1])], days=4, weights=SPACING)
    assert len(notes(even)) == 1 and "(2 of 4)" in notes(even)[0].message  # a 4-day week keeps only 2 apart
