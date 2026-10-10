import io

import pytest
from openpyxl import Workbook

from automated_academics.checks import diagnose
from automated_academics.grid_io import GridError, GridOptions, import_grid
from automated_academics.solver import solve
from automated_academics.validate import find_conflicts

DAYS = ["MON", "TUE", "WED", "THU", "FRI", "SAT"]
TIMES = ["8 TO 8.50", "8.55 TO 9.45", "9.50 TO 10.40", "11 TO 11.50"]


def put(ws, top, title, rows, times=TIMES, days=DAYS, subtitle="FROM 3RD JANUARY 2022", col=1):
    """A block like a college would draw it: heading, a 'from' line, the day row, then one row per lecture."""
    ws.cell(row=top, column=col, value=title)
    if subtitle:
        ws.cell(row=top + 1, column=col + 2, value=subtitle)
    for j, d in enumerate(days):
        ws.cell(row=top + 2, column=col + 1 + j, value=d)
    for i, row in enumerate(rows):
        ws.cell(row=top + 3 + i, column=col, value=times[i] if i < len(times) else None)
        for j, v in enumerate(row):
            if v:
                ws.cell(row=top + 3 + i, column=col + 1 + j, value=v)
    return top + 3 + len(rows)


def book(build) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def one_class():
    def build(wb):
        ws = wb.create_sheet("Sem 1")
        put(ws, 1, "BSc Maths Sem 1 ( A BLOCK, 1ST FLOOR, ROOM NO_101)", [
            ["Calculus(Prof. Anil Rao)", "Algebra(Dr. Meena Iyer)", "Calculus(Prof. Anil Rao)", "", "Algebra(Dr. Meena Iyer)", "Statistics(Prof. Ravi Nair)"],
            ["Algebra(Dr. Meena Iyer)", "Calculus(Prof. Anil Rao)", "Statistics(Prof. Ravi Nair)", "Calculus(Prof. Anil Rao)", "", ""],
            ["", "Statistics(Prof. Ravi Nair)", "", "Algebra(Dr. Meena Iyer)", "Calculus(Prof. Anil Rao)", "Algebra(Dr. Meena Iyer)"],
            ["Statistics(Prof. Ravi Nair)", "", "Algebra(Dr. Meena Iyer)", "Statistics(Prof. Ravi Nair)", "Statistics(Prof. Ravi Nair)", "Calculus(Prof. Anil Rao)"],
        ])
    return book(build)


def by_name(inst):
    return {c.name: c for c in inst.courses}, {f.name: f for f in inst.faculty}


# ---------- the basics ----------

def test_reads_a_simple_grid_into_an_institution():
    inst, report, placements = import_grid(one_class())
    assert [b.program for b in inst.batches] == ["BSc Maths Sem 1"]
    assert inst.batches[0].semester == 1
    courses, faculty = by_name(inst)
    assert set(courses) == {"Calculus", "Algebra", "Statistics"}
    assert set(faculty) == {"Prof. Anil Rao", "Dr. Meena Iyer", "Prof. Ravi Nair"}
    assert inst.calendar.day_names == ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
    assert inst.calendar.lectures_per_day == 4
    assert sum(len(o.sessions) for o in inst.offerings) == len(placements) == sum(
        1 for _ in range(1) for _ in [0]) * 0 + len(placements)
    assert [r.name for r in inst.rooms] == ["Room 101"]
    assert report.existing["clashes"] == 0 and report.existing["sessions"] == len(placements)


def test_the_times_down_the_side_become_lecture_times_and_a_break():
    inst, report, _ = import_grid(one_class())
    t = inst.calendar.times
    assert [(x.start, x.end) for x in t] == [("08:00", "08:50"), ("08:55", "09:45"), ("09:50", "10:40"), ("11:00", "11:50")]
    assert inst.calendar.break_after == [2]  # the 20-minute gap after the third lecture
    assert any("break after lecture 3" in a for a in report.assumptions)


def test_a_row_marked_as_a_break_is_a_break_not_a_lecture():
    def build(wb):
        ws = wb.create_sheet("W")
        top = put(ws, 1, "BA History Sem 2", [["H(Prof. A One)", "H(Prof. A One)", "G(Prof. B Two)"], ["G(Prof. B Two)", "", "H(Prof. A One)"]],
                  times=["9 TO 10", "10 TO 11"], days=DAYS[:3])
        ws.cell(row=top, column=1, value="LUNCH")
        ws.cell(row=top + 1, column=1, value="12 TO 1")
        ws.cell(row=top + 1, column=2, value="G(Prof. B Two)")
    inst, report, _ = import_grid(book(build))
    assert inst.calendar.lectures_per_day == 3
    assert inst.calendar.break_after == [1]


def test_nothing_that_looks_like_a_grid_is_an_error():
    def build(wb):
        wb.create_sheet("Notes").append(["just", "some", "text"])
    with pytest.raises(GridError):
        import_grid(book(build))


# ---------- cells ----------

def test_the_usual_ways_of_writing_a_subject_and_teacher():
    def build(wb):
        ws = wb.create_sheet("S")
        put(ws, 1, "MBA Sem 3", [
            ["Finance\nProf. Kiran Das", "Marketing - Dr. Leela Rao", "Strategy Prof. Omar Khan", "Law (Room 3)"],
            ["Finance (Prof. Kiran Das)", "", "", ""],
        ], times=["9 TO 10", "10 TO 11"], days=DAYS[:4])
    inst, report, _ = import_grid(book(build))
    courses, faculty = by_name(inst)
    assert {"Finance", "Marketing", "Strategy", "Law (Room 3)"} == set(courses)
    assert set(faculty) >= {"Prof. Kiran Das", "Dr. Leela Rao", "Prof. Omar Khan"}
    # a bracket that is a room, not a person, is not taken as a teacher
    assert not any("Room" in f.name for f in inst.faculty if f.name.startswith("Room"))
    assert any("teacher not given" in f.name for f in inst.faculty)
    assert any("no teacher" in p for p in report.problems)


def test_spellings_of_one_person_are_merged_and_reported():
    def build(wb):
        ws = wb.create_sheet("S")
        put(ws, 1, "BCom Sem 1", [
            ["Tax(Prof. Anil Rao)", "Tax(Anil Rao)", "Audit(Prof Anil Rao)", "Law(Dr. Seema Pillai)"],
            ["Law(Seema Pillai)", "Audit(prof. anil rao)", "", ""],
        ], times=["9 TO 10", "10 TO 11"], days=DAYS[:4])
    inst, report, _ = import_grid(book(build))
    assert len(inst.faculty) == 2
    assert any(m.startswith("teacher:") and "Anil Rao" in m for m in report.merged)


def test_a_bare_surname_that_fits_two_people_is_not_guessed():
    def build(wb):
        ws = wb.create_sheet("S")
        put(ws, 1, "BBA Sem 1", [["A(Prof. Amit Shah)", "B(Prof. Dhruv Shah)", "C(Shah)", ""]], times=["9 TO 10"], days=DAYS[:4])
    inst, report, _ = import_grid(book(build))
    assert len(inst.faculty) == 3
    assert any("'Shah' could be" in p for p in report.problems)


def test_numbered_options_are_parallel_electives_in_groups():
    def build(wb):
        ws = wb.create_sheet("S")
        put(ws, 1, "MBA Sem 4 (Room 5)", [
            ["Core(Prof. A One)", "1) Fin Risk (Prof. B Two) 2) Brand Mgmt (Prof. C Three)", "Core(Prof. A One)"],
            ["Core(Prof. A One)", "1) Fin Risk (Prof. B Two) 2) Brand Mgmt (Prof. C Three)", ""],
        ], times=["9 TO 10", "10 TO 11"], days=DAYS[:3])
    inst, report, placements = import_grid(book(build))
    groups = [b for b in inst.batches if b.group_of]
    assert len(groups) == 2 and {g.group_of for g in groups} == {inst.batches[0].id}
    fin = next(o for o in inst.offerings if inst.courses and next(c for c in inst.courses if c.code == o.course_code).name == "Fin Risk")
    assert fin.batch_ids[0] in {g.id for g in groups}
    assert len(report.parallel) == 2
    assert any("spare room" in r.name.lower() for r in inst.rooms)
    assert find_conflicts(inst, __import__("automated_academics.models", fromlist=["Timetable"]).Timetable(placements=placements, status="x")) == []


def test_a_faculty_table_beside_the_grid_names_the_teacher_of_a_bare_cell():
    def build(wb):
        ws = wb.create_sheet("S")
        put(ws, 1, "BSc Phys Sem 1", [["Optics", "Optics", "Mechanics(Prof. Zed Zee)"]], times=["9 TO 10"], days=DAYS[:3])
        ws.cell(row=1, column=9, value="Subject Name")
        ws.cell(row=1, column=10, value="Faculty Name")
        ws.cell(row=2, column=9, value="Optics")
        ws.cell(row=2, column=10, value="Dr. Opal Sen")
    inst, report, _ = import_grid(book(build))
    optics = next(c for c in inst.courses if c.name == "Optics")
    off = next(o for o in inst.offerings if o.course_code == optics.code)
    assert next(f for f in inst.faculty if f.id == off.faculty_id).name == "Dr. Opal Sen"
    assert not any("teacher not given" in f.name for f in inst.faculty)


# ---------- sessions ----------

def double_lecture_grid():
    def build(wb):
        ws = wb.create_sheet("S")
        put(ws, 1, "BA Eng Sem 1", [
            ["Poetry(Prof. A One)", "", ""],
            ["Poetry(Prof. A One)", "Prose(Prof. B Two)", ""],
            ["Prose(Prof. B Two)", "Prose(Prof. B Two)", "Poetry(Prof. A One)"],
            ["Poetry(Prof. A One)", "", ""],
        ], days=DAYS[:3])
    return book(build)


def test_consecutive_lectures_become_a_double_lecture_but_not_across_the_break():
    inst, _, placements = import_grid(double_lecture_grid())
    poetry = next(c for c in inst.courses if c.name == "Poetry")
    off = next(o for o in inst.offerings if o.course_code == poetry.code)
    # Mon L1+L2 are one block of 2; Mon L4 is separate (a break falls after L3); Wed L3 stands alone
    assert sorted(off.sessions) == [1, 1, 2]
    assert inst.calendar.break_after == [2]
    assert sum(p.length for p in placements) == 7  # every filled cell is accounted for


def test_doubles_can_be_switched_off():
    inst, _, _ = import_grid(double_lecture_grid(), GridOptions(merge_consecutive=False))
    assert all(s == 1 for o in inst.offerings for s in o.sessions)


# ---------- several sheets ----------

def test_a_class_repeated_on_another_sheet_is_read_once_and_a_teachers_own_view_is_skipped():
    def build(wb):
        a = wb.create_sheet("Week 1")
        put(a, 1, "BBA A", [["M(Prof. A One)", "N(Prof. B Two)", ""]], times=["9 TO 10"], days=DAYS[:3])
        b = wb.create_sheet("Week 2 (short)")
        put(b, 1, "BBA A", [["M(Prof. A One)", "N(Prof. B Two)", ""]], times=["9 TO 9.40"], days=DAYS[:3])
        c = wb.create_sheet("Individual Time Table")
        put(c, 1, "Prof. A One", [["BBA A", "", ""]], times=["9 TO 10"], days=DAYS[:3], subtitle=None)
    inst, report, _ = import_grid(book(build))
    assert [b.program for b in inst.batches] == ["BBA A"]
    notes = {s["name"]: s["note"] for s in report.sheets}
    assert "already read" in notes["Week 2 (short)"]
    assert "one teacher's own timetable" in notes["Individual Time Table"]
    assert len(inst.offerings) == 2


def test_sheets_can_be_chosen():
    def build(wb):
        put(wb.create_sheet("A"), 1, "Class A", [["M(Prof. A One)", "", ""]], times=["9 TO 10"], days=DAYS[:3])
        put(wb.create_sheet("B"), 1, "Class B", [["N(Prof. B Two)", "", ""]], times=["9 TO 10"], days=DAYS[:3])
    inst, report, _ = import_grid(book(build), GridOptions(sheets=["B"]))
    assert [b.program for b in inst.batches] == ["Class B"]
    assert [s["name"] for s in report.sheets] == ["B"]


# ---------- the existing timetable, scored ----------

def test_clashes_and_idle_gaps_in_the_existing_timetable_are_counted():
    def build(wb):
        ws = wb.create_sheet("S")
        put(ws, 1, "Class A", [["Math(Prof. A One)", "", ""], ["", "", ""], ["Phy(Prof. B Two)", "", ""]], times=["9 TO 10", "10 TO 11", "11 TO 12"], days=DAYS[:3])
        put(ws, 8, "Class B", [["Math(Prof. A One)", "", ""]], times=["9 TO 10"], days=DAYS[:3])  # same teacher, same slot
    inst, report, _ = import_grid(book(build))
    assert report.existing["clashes"] >= 1
    assert report.existing["quality"]["batch_gaps"] == 1  # Class A is idle in the middle of Monday
    assert any(c.startswith("clash F:") for c in report.existing["clash_examples"])  # the teacher is in two places


# ---------- what comes out works ----------

def test_the_imported_institution_is_valid_and_solvable():
    inst, report, _ = import_grid(one_class())
    assert [d for d in diagnose(inst) if d.level == "error"] == []
    tt = solve(inst, time_limit_s=5)
    assert len(tt.placements) == sum(len(o.sessions) for o in inst.offerings)
    assert find_conflicts(inst, tt) == []


def test_assumptions_are_stated():
    _, report, _ = import_grid(one_class())
    text = " ".join(report.assumptions)
    assert "60 students" in text and "assumed to seat" in text


# ---------- the API ----------

def test_the_endpoint_previews_without_saving_and_the_timetable_can_be_kept(tmp_path):
    from fastapi.testclient import TestClient
    from automated_academics.api import create_app
    with TestClient(create_app(str(tmp_path / "t.db"))) as c:
        r = c.post("/institutions/import-grid", files={"file": ("grid.xlsx", one_class().getvalue())})
        assert r.status_code == 200, r.text
        body = r.json()
        assert c.get("/institutions").json() == []  # a preview saves nothing
        assert body["report"]["existing"]["sessions"] == len(body["placements"]) > 0
        assert not any(i["level"] == "error" for i in body["issues"])

        iid = c.post("/institutions", json=body["institution"]).json()["id"]
        kept = c.post(f"/institutions/{iid}/timetable/imported", json={"placements": body["placements"]})
        assert kept.status_code == 201 and kept.json()["report"]["ok"] is True
        latest = c.get(f"/institutions/{iid}/latest-job").json()
        assert latest["id"] == kept.json()["job_id"]
        tt = c.get(f"/jobs/{latest['id']}/timetable").json()
        assert tt["status"] == "IMPORTED" and len(tt["placements"]) == len(body["placements"])

        bad = dict(body["placements"][0], offering_id="GHOST")
        assert c.post(f"/institutions/{iid}/timetable/imported", json={"placements": [bad]}).status_code == 422


def test_the_endpoint_rejects_what_is_not_a_grid_or_a_workbook(tmp_path):
    from fastapi.testclient import TestClient
    from automated_academics.api import create_app
    with TestClient(create_app(str(tmp_path / "t.db"))) as c:
        def build(wb):
            wb.create_sheet("Notes").append(["nothing", "here"])
        assert c.post("/institutions/import-grid", files={"file": ("n.xlsx", book(build).getvalue())}).status_code == 422
        assert c.post("/institutions/import-grid", files={"file": ("x.xlsx", b"not a workbook")}).status_code == 400
        assert c.post("/institutions/import-grid", files={"file": ("g.xlsx", one_class().getvalue())},
                      data={"options": '{"default_students": 0}'}).status_code == 422


# ---------- guards against over-eager reading ----------

def test_a_stray_day_name_is_not_a_timetable_block():
    def build(wb):
        ws = wb.create_sheet("S")
        ws.cell(row=1, column=1, value="Note")
        ws.cell(row=1, column=2, value="Monday")
        ws.cell(row=1, column=3, value="Friday")
        ws.cell(row=2, column=1, value="9 TO 10")
        ws.cell(row=2, column=2, value="Meeting(Prof. Q Zed)")
        put(ws, 5, "Class A", [["M(Prof. A One)", "", ""]], times=["9 TO 10"], days=DAYS[:3])
    inst, _, _ = import_grid(book(build))
    assert [b.program for b in inst.batches] == ["Class A"]


def test_a_block_with_no_heading_does_not_borrow_one_from_the_block_above():
    def build(wb):
        ws = wb.create_sheet("S")
        end = put(ws, 1, "Class A", [["Alpha(Prof. A One)", "", ""], ["Beta(Prof. B Two)", "", ""]], times=["9 TO 10", "10 TO 11"], days=DAYS[:3])
        ws.cell(row=end, column=1, value="")
        # no heading: the day row follows straight after Class A's last lecture row
        for j, d in enumerate(DAYS[:3]):
            ws.cell(row=end + 1, column=2 + j, value=d)
        ws.cell(row=end + 2, column=1, value="9 TO 10")
        ws.cell(row=end + 2, column=2, value="Gamma(Prof. C Three)")
    inst, _, _ = import_grid(book(build))
    assert [b.program for b in inst.batches] == ["Class A", "Class 2 (sheet S)"]


def test_the_same_class_under_a_different_heading_is_still_recognised_as_a_repeat():
    def build(wb):
        a = wb.create_sheet("Week 1")
        rows = [["M(Prof. A One)", "N(Prof. B Two)", "O(Prof. C Three)"], ["N(Prof. B Two)", "M(Prof. A One)", ""]]
        put(a, 1, "BBA Sem 1", rows, times=["9 TO 10", "10 TO 11"], days=DAYS[:3])
        b = wb.create_sheet("Week 2")
        put(b, 1, "Bachelor of Business (First semester)", rows, times=["9 TO 9.40", "9.45 TO 10.25"], days=DAYS[:3])
    inst, report, _ = import_grid(book(build))
    assert len(inst.batches) == 1
    assert "same subjects in the same places" in {s["name"]: s["note"] for s in report.sheets}["Week 2"]


def test_a_teacher_wise_sheet_without_a_person_heading_is_still_skipped():
    def build(wb):
        put(wb.create_sheet("Classes"), 1, "Class A", [["M(Prof. A One)", "", ""]], times=["9 TO 10"], days=DAYS[:3])
        ws = wb.create_sheet("Faculty timetable")
        ws.cell(row=1, column=2, value="MON")
        ws.cell(row=1, column=3, value="TUE")
        ws.cell(row=1, column=4, value="WED")
        ws.cell(row=2, column=1, value="9 TO 10")
        ws.cell(row=2, column=2, value="Class A")
    inst, report, _ = import_grid(book(build))
    assert [b.program for b in inst.batches] == ["Class A"]
    assert "teacher" in {s["name"]: s["note"] for s in report.sheets}["Faculty timetable"]


def test_a_shortened_surname_is_the_same_person():
    def build(wb):
        ws = wb.create_sheet("S")
        put(ws, 1, "Class A", [["M(Prof. Chander Rajpurohit)", "N(Prof. Chander Raj)", "O(Chander)"]], times=["9 TO 10"], days=DAYS[:3])
    inst, report, _ = import_grid(book(build))
    assert len(inst.faculty) == 1
    assert not any("could be" in p for p in report.problems)


def test_a_double_lecture_never_runs_across_a_break():
    def build(wb):
        ws = wb.create_sheet("S")
        put(ws, 1, "Class A", [[""] * 3, [""] * 3, ["Law(Prof. A One)", "", ""], ["Law(Prof. A One)", "", ""]], days=DAYS[:3])
    inst, _, _ = import_grid(book(build))
    assert inst.calendar.break_after == [2]
    law = inst.offerings[0]
    assert law.sessions == [1, 1]


def test_a_small_difference_in_gaps_is_not_a_break():
    def build(wb):
        ws = wb.create_sheet("S")
        # gaps of 5, 5 and 10 minutes: the 10 is twice the usual but still just a changeover
        put(ws, 1, "Class A", [["M(Prof. A One)", "", ""]] * 4, times=["9 TO 9.50", "9.55 TO 10.45", "10.50 TO 11.40", "11.50 TO 12.40"], days=DAYS[:3])
    inst, report, _ = import_grid(book(build))
    assert inst.calendar.break_after == []
    assert any("No break was found" in a for a in report.assumptions)
