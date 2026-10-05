import pytest
from openpyxl import load_workbook

from automated_academics.cli import main
from automated_academics.excel_io import ImportErrors, export_workbook, import_workbook
from automated_academics.solver import solve
from automated_academics.synthetic import sample_institution
from automated_academics.validate import find_conflicts


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "inst.xlsx"
    export_workbook(sample_institution(), path)
    return path


def edit(path, sheet, row, col_name, value):
    wb = load_workbook(path)
    ws = wb[sheet]
    col = next(i for i, c in enumerate(ws[1], start=1) if c.value == col_name)
    ws.cell(row=row, column=col, value=value)
    wb.save(path)


def test_round_trip_preserves_model(workbook):
    original = sample_institution()
    assert import_workbook(workbook).model_dump() == original.model_dump()


def test_imported_workbook_solves(workbook):
    inst = import_workbook(workbook)
    tt = solve(inst, time_limit_s=4)
    assert find_conflicts(inst, tt) == []


def test_all_problems_reported_together(workbook):
    edit(workbook, "Rooms", 2, "capacity", "lots")
    edit(workbook, "Batches", 3, "level", "Diploma")
    edit(workbook, "Faculty", 2, "unavailable", "Funday:1")
    with pytest.raises(ImportErrors) as exc:
        import_workbook(workbook)
    text = str(exc.value)
    assert "Rooms row 2" in text and "capacity" in text
    assert "Batches row 3" in text and "level" in text
    assert "Faculty row 2" in text and "Mon:1" in text
    assert len(exc.value.issues) == 3


def test_unknown_reference_is_reported(workbook):
    edit(workbook, "Offerings", 2, "faculty_id", "NOBODY")
    with pytest.raises(ImportErrors) as exc:
        import_workbook(workbook)
    assert "NOBODY" in str(exc.value)


def test_duplicate_ids_are_reported(workbook):
    edit(workbook, "Rooms", 3, "id", "C1")
    with pytest.raises(ImportErrors) as exc:
        import_workbook(workbook)
    assert "duplicate id 'C1'" in str(exc.value)


def test_missing_sheet_and_column(workbook):
    wb = load_workbook(workbook)
    del wb["Rooms"]
    wb["Faculty"]["C1"] = "dept"  # rename required column 'department'
    wb.save(workbook)
    with pytest.raises(ImportErrors) as exc:
        import_workbook(workbook)
    text = str(exc.value)
    assert "Rooms: sheet not found" in text


def institution_row(path, key):
    ws = load_workbook(path)["Institution"]
    return next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == key)


def test_weights_and_avoid_round_trip(tmp_path):
    inst = sample_institution()
    inst.weights = inst.weights.model_copy(update={"batch_gaps": 25, "avoid_slot": 0})
    p = tmp_path / "w.xlsx"
    export_workbook(inst, p)
    back = import_workbook(p)
    assert back.weights == inst.weights
    assert back.faculty[0].avoid == inst.faculty[0].avoid and back.faculty[0].avoid  # not vacuous


def test_blank_weight_uses_default_and_zero_is_kept(workbook):
    wb = load_workbook(workbook)
    ws = wb["Institution"]
    ws.cell(row=institution_row(workbook, "weight_batch_gaps"), column=2, value=None)
    ws.cell(row=institution_row(workbook, "weight_faculty_gaps"), column=2, value=0)
    wb.save(workbook)
    w = import_workbook(workbook).weights
    assert w.batch_gaps == 10 and w.faculty_gaps == 0


@pytest.mark.parametrize("bad", [-1, "lots", 2.5])
def test_bad_weight_is_reported(workbook, bad):
    wb = load_workbook(workbook)
    wb["Institution"].cell(row=institution_row(workbook, "weight_peak_day_load"), column=2, value=bad)
    wb.save(workbook)
    with pytest.raises(ImportErrors) as exc:
        import_workbook(workbook)
    assert "weight" in str(exc.value) and exc.value.issues[0].sheet == "Institution"


def test_bad_avoid_entry_is_reported(workbook):
    edit(workbook, "Faculty", 2, "avoid", "Someday:3")
    with pytest.raises(ImportErrors) as exc:
        import_workbook(workbook)
    assert "Faculty row 2" in str(exc.value) and "avoid" in str(exc.value)


def test_workbook_without_avoid_column_or_weights_still_imports(workbook):
    wb = load_workbook(workbook)
    fac = wb["Faculty"]
    col = next(i for i, c in enumerate(fac[1], start=1) if c.value == "avoid")
    fac.delete_cols(col)
    inst_ws = wb["Institution"]
    for r in sorted((r for r in range(2, inst_ws.max_row + 1)
                     if str(inst_ws.cell(row=r, column=1).value).startswith("weight_")), reverse=True):
        inst_ws.delete_rows(r)
    wb.save(workbook)
    inst = import_workbook(workbook)
    assert all(f.avoid == [] for f in inst.faculty)
    assert inst.weights.batch_gaps == 10  # defaults


def test_cli_template_then_check(tmp_path, capsys):
    out = tmp_path / "t.xlsx"
    assert main(["template", str(out)]) == 0
    assert main(["check", str(out)]) == 0
    assert "OK:" in capsys.readouterr().out
    edit(out, "Rooms", 2, "capacity", "x")
    assert main(["check", str(out)]) == 1
