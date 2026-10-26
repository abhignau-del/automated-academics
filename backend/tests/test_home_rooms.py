"""Named home rooms: a class's classroom lectures go in its own room."""

import io

import pytest
from pydantic import ValidationError

from automated_academics.checks import diagnose
from automated_academics.decompose import room_pool
from automated_academics.excel_io import export_workbook, import_workbook
from automated_academics.models import Institution, Pin, RoomKind
from automated_academics.solver import solve
from automated_academics.synthetic import sample_institution
from automated_academics.validate import find_conflict_details, find_conflicts
from collections import Counter


def homed(**homes):
    """The sample with home rooms given as batch id -> room id."""
    inst = sample_institution()
    for b in inst.batches:
        if b.id in homes:
            b.home_room = homes[b.id]
    Institution.model_validate(inst.model_dump())  # still valid
    return inst


def classroom_rows(inst, tt):
    """(placement, offering, course) for every lecture held in a classroom-kind course."""
    offs = {o.id: o for o in inst.offerings}
    courses = {c.code: c for c in inst.courses}
    return [(p, offs[p.offering_id], courses[offs[p.offering_id].course_code]) for p in tt.placements
            if courses[offs[p.offering_id].course_code].room_kind is RoomKind.CLASSROOM]


def errors(inst):
    return [d.message for d in diagnose(inst) if d.level == "error"]


# ---------- the model ----------

def test_a_home_room_must_exist():
    inst = sample_institution()
    inst.batches[0].home_room = "nowhere"
    with pytest.raises(ValidationError, match="unknown home room"):
        Institution.model_validate(inst.model_dump())


def test_a_sub_group_does_not_inherit_its_parents_home_room_but_may_have_its_own():
    inst = homed(**{"CS-UG1": "C1"})
    assert inst.home_room_of_batch("CS-UG1") == "C1"
    assert inst.home_room_of_batch("CS-UG1-A") is None and inst.home_room_of_batch("CS-UG3") is None
    next(b for b in inst.batches if b.id == "CS-UG1-B").home_room = "C2"
    assert inst.home_room_of_batch("CS-UG1-B") == "C2"
    assert inst.home_room_of_batch("no-such-class") is None


def test_classes_attending_together_follow_the_one_home_room_they_agree_on():
    inst = homed(**{"CS-UG1": "C1", "MG-UG1": "C2"})
    assert inst.home_room_of(["CS-UG1"]) == "C1"
    assert inst.home_room_of(["CS-UG1", "CS-UG3"]) == "C1"  # CS-UG3 has none, so goes along
    assert inst.home_room_of(["CS-UG1", "MG-UG1"]) is None  # they disagree: neither can be at home
    assert inst.home_room_of(["CS-UG3", "CS-PG1"]) is None


# ---------- the solver ----------

def test_classroom_lectures_stay_in_the_home_room_and_labs_are_free():
    inst = homed(**{"CS-UG1": "C1", "MG-UG1": "C2", "CS-PG1": "C3"})
    tt = solve(inst, time_limit_s=10)
    assert find_conflicts(inst, tt) == []
    checked = 0
    for p, off, course in classroom_rows(inst, tt):
        home = inst.home_room_of(off.batch_ids)
        if home:
            assert p.room_id == home, (off.id, p.room_id, home)
            checked += 1
    assert checked > 0
    assert all(p.room_id in {"L1", "L2"} for p in tt.placements
               if next(c for c in inst.courses if c.code == next(o for o in inst.offerings if o.id == p.offering_id).course_code).room_kind is RoomKind.LAB)


def test_the_home_room_holds_for_whole_class_lectures_but_not_for_its_practical_groups():
    inst = homed(**{"CS-UG1": "C4"})
    tt = solve(inst, time_limit_s=10)
    whole = [p for p, off, _ in classroom_rows(inst, tt) if off.batch_ids == ["CS-UG1"]]
    assert whole and all(p.room_id == "C4" for p in whole)
    groups = [off for _, off, _ in classroom_rows(inst, tt) if set(off.batch_ids) <= {"CS-UG1-A", "CS-UG1-B"}]
    assert all(inst.home_room_of(off.batch_ids) is None for off in groups)  # free to meet in parallel elsewhere


def test_no_home_rooms_means_nothing_changes():
    inst = sample_institution()
    assert all(b.home_room is None for b in inst.batches)
    tt = solve(inst, time_limit_s=6)
    assert find_conflicts(inst, tt) == []


# ---------- the checks ----------

def test_a_home_room_too_small_for_the_class_is_an_error():
    inst = homed(**{"CS-UG1": "C1"})
    inst.rooms[0].capacity = 10
    msgs = [m for m in errors(inst) if "lives in" in m]
    assert msgs and "Classroom 1" in msgs[0] and "10 seats" in msgs[0]


def test_classes_that_together_need_more_than_their_shared_room_has_are_an_error():
    inst = homed(**{b.id: "C1" for b in sample_institution().batches})
    msgs = [m for m in errors(inst) if "can only use Classroom 1" in m]
    assert msgs


def test_a_joint_lecture_of_classes_with_different_homes_is_a_note_not_an_error():
    inst = homed(**{"CS-UG1": "C1", "MG-UG1": "C2"})
    from automated_academics.models import Offering
    course = next(c for c in inst.courses if c.room_kind is RoomKind.CLASSROOM)
    fac = inst.faculty[0]
    inst.offerings.append(Offering(id="JOINT", course_code=course.code, faculty_id=fac.id, batch_ids=["CS-UG1", "MG-UG1"], sessions=[1]))
    notes = [d for d in diagnose(inst) if "different rooms" in d.message]
    assert len(notes) == 1 and notes[0].level == "info" and "JOINT" in notes[0].message
    assert not [m for m in errors(inst) if "JOINT" in m]


def test_a_pin_to_another_room_than_home_is_reported_and_the_home_room_is_accepted():
    inst = homed(**{"CS-UG1": "C1"})
    off = next(o for o in inst.offerings if o.batch_ids == ["CS-UG1"]
               and next(c for c in inst.courses if c.code == o.course_code).room_kind is RoomKind.CLASSROOM)
    inst.pins = [Pin(offering_id=off.id, session_index=0, day=0, start=0, room_id="C2")]
    assert any("lives in Classroom 1" in m for m in errors(inst))
    inst.pins = [Pin(offering_id=off.id, session_index=0, day=0, start=0, room_id="C1")]
    assert not any("lives in" in m for m in errors(inst))


# ---------- manual edits ----------

def test_moving_a_lecture_out_of_the_home_room_is_a_conflict_and_putting_it_back_clears_it():
    inst = homed(**{"CS-UG1": "C1"})
    tt = solve(inst, time_limit_s=8)
    p, off, _ = next(x for x in classroom_rows(inst, tt) if inst.home_room_of(x[1].batch_ids) == "C1")
    moved = tt.model_copy(deep=True)
    target = next(q for q in moved.placements if (q.offering_id, q.session_index) == (p.offering_id, p.session_index))
    target.room_id = "C2"
    hits = [c for c in find_conflict_details(inst, moved) if c.kind == "home-room"]
    assert len(hits) == 1 and "Classroom 1" in hits[0].message and "Classroom 2" in hits[0].message
    target.room_id = "C1"
    assert not [c for c in find_conflict_details(inst, moved) if c.kind == "home-room"]


# ---------- the decomposed solver ----------

def test_a_small_room_pool_for_a_group_always_includes_the_home_room():
    inst = homed(**{"CS-UG1": "C4"})
    ids = [o.id for o in inst.offerings if o.batch_ids == ["CS-UG1"]
           and next(c for c in inst.courses if c.code == o.course_code).room_kind is RoomKind.CLASSROOM]
    assert ids
    assert "C4" in room_pool(inst, ids, Counter({"C4": 99}))  # even if C4 looks busy and others are free


# ---------- the workbook ----------

def test_home_rooms_survive_the_workbook_by_id_and_by_name():
    import openpyxl
    from automated_academics.excel_io import ImportErrors
    inst = homed(**{"CS-UG1": "C1", "MG-UG1": "C2"})
    buf = io.BytesIO()
    export_workbook(inst, buf)
    back = import_workbook(io.BytesIO(buf.getvalue()))
    assert back.home_room_of_batch("CS-UG1") == "C1" and back.home_room_of_batch("MG-UG1") == "C2"
    wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
    ws = wb["Batches"]
    col = [c.value for c in ws[1]].index("home_room") + 1
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "CS-UG3")
    ws.cell(row, col).value = "classroom 3"  # written by name, any capitals
    out = io.BytesIO()
    wb.save(out)
    assert import_workbook(io.BytesIO(out.getvalue())).home_room_of_batch("CS-UG3") == "C3"
    ws.cell(row, col).value = "Room 99"
    bad = io.BytesIO()
    wb.save(bad)
    with pytest.raises(ImportErrors) as e:
        import_workbook(io.BytesIO(bad.getvalue()))
    assert "home room" in str(e.value.issues)
