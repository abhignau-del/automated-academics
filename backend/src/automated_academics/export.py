"""Excel and PDF export of a timetable.

Layout rules shared by both formats
  * one week grid (lectures x days) per class, faculty member and room
  * a multi-lecture session is merged vertically unless it shares cells with a
    parallel session (e.g. two lab groups at once); then each cell lists the
    sessions compactly instead
"""

from __future__ import annotations

import colorsys
import logging
import re
from typing import Any, BinaryIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Calendar, Institution, LectureTime, Timetable
from .views import KINDS, Kind, entities, session_views

log = logging.getLogger("automated_academics")

KIND_TITLES = {"batch": "Class", "faculty": "Faculty", "room": "Room"}
Session = dict[str, Any]


# ---------- shared layout ----------

def course_rgb(code: str) -> tuple[float, float, float]:
    """Same pastel per course code as the web UI (hue hashed from the code)."""
    h = 0
    for ch in code:
        h = (h * 31 + ord(ch)) % 360
    return colorsys.hls_to_rgb(h / 360, 0.88, 0.70)


def course_hex(code: str) -> str:
    r, g, b = (round(c * 255) for c in course_rgb(code))
    return f"{r:02X}{g:02X}{b:02X}"


def layout(sessions: list[Session]) -> tuple[dict[tuple[int, int], list[int]], list[int]]:
    """Map each (day, lecture) cell to the sessions covering it.

    Returns (cover, merge) where `merge` lists indices of multi-lecture sessions
    that own all of their cells exclusively and can therefore be merged.
    """
    cover: dict[tuple[int, int], list[int]] = {}
    for i, s in enumerate(sessions):
        for t in range(s["start"], s["start"] + s["length"]):
            cover.setdefault((s["day"], t), []).append(i)
    merge = [
        i for i, s in enumerate(sessions)
        if s["length"] > 1 and all(cover[(s["day"], t)] == [i]
                                   for t in range(s["start"], s["start"] + s["length"]))
    ]
    return cover, merge


def full_lines(kind: Kind, s: Session) -> list[str]:
    head = f"{s['course_code']} {s['course_name']}"
    batches = ", ".join(s["batch_ids"])
    if kind == "batch":
        return [head, s["faculty_name"], s["room_name"]]
    if kind == "faculty":
        return [head, batches, s["room_name"]]
    return [head, s["faculty_name"], batches]


def compact_line(kind: Kind, s: Session) -> str:
    extra = {"batch": [s["faculty_name"], s["room_name"]],
             "faculty": [", ".join(s["batch_ids"]), s["room_name"]],
             "room": [s["faculty_name"], ", ".join(s["batch_ids"])]}[kind]
    return " · ".join([s["course_code"], *extra])


def lecture_label(s: Session) -> str:
    return f"L{s['start'] + 1}" if s["length"] == 1 else f"L{s['start'] + 1}–L{s['start'] + s['length']}"


def _selection(inst: Institution, tt: Timetable, kind: str | None, ident: str | None):
    """Entities to export as (kind, id, title, sessions). Empty entities are skipped
    unless explicitly requested. Raises LookupError for an unknown id."""
    kinds = [kind] if kind else list(KINDS)
    out = []
    for k in kinds:
        for eid, label in entities(inst, k):  # type: ignore[arg-type]
            if ident is not None and eid != ident:
                continue
            sessions = session_views(inst, tt, k, eid)  # type: ignore[arg-type]
            if sessions or ident is not None:
                out.append((k, eid, label, sessions))
    if ident is not None and not out:
        raise LookupError(f"{kind} {ident!r} not found")
    return out


def _cell_text(kind: Kind, sessions: list[Session], idxs: list[int], first: bool) -> list[str]:
    if len(idxs) == 1:
        s = sessions[idxs[0]]
        return full_lines(kind, s) if first else [f"↓ {s['course_code']}"]
    return [compact_line(kind, sessions[i]) for i in idxs]


# ---------- Excel ----------

_THIN = Side(style="thin", color="999999")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(prefix: str, ident: str, taken: set[str]) -> str:
    base = _FORBIDDEN.sub("-", f"{prefix} {ident}")[:31]
    name, n = base, 2
    while name.lower() in taken:
        suffix = f" ({n})"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    taken.add(name.lower())
    return name


def _write_grid(ws, cal: Calendar, kind: Kind, title: str, sessions: list[Session]) -> None:
    ncols = cal.days + 1
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)

    ws.cell(row=2, column=1, value="")
    for d, name in enumerate(cal.day_names):
        c = ws.cell(row=2, column=d + 2, value=name)
        c.fill, c.font, c.border = _HEAD_FILL, _HEAD_FONT, _BOX
        c.alignment = Alignment(horizontal="center")

    cover, merge = layout(sessions)
    merged = set(merge)
    starts = {(s["day"], s["start"]): i for i, s in enumerate(sessions)}

    for p in range(cal.lectures_per_day):
        row = p + 3
        shown = cal.times[p] if cal.times else None
        label = ws.cell(row=row, column=1, value=f"L{p + 1}" + (
            f"\n{LectureTime.pretty(shown.start)}\u2013{LectureTime.pretty(shown.end)}" if shown else ""))
        label.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        label.font, label.border = Font(bold=True), _BOX
        label.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 62
        for d in range(cal.days):
            idxs = cover.get((d, p), [])
            cell = ws.cell(row=row, column=d + 2)
            cell.border = _BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if not idxs:
                continue
            owner = idxs[0] if len(idxs) == 1 else None
            first = owner is not None and starts.get((d, p)) == owner
            if owner in merged and not first:
                continue  # covered by a merged range
            lines = _cell_text(kind, sessions, idxs, first)
            if first and d in cal.day_times:  # this day runs to its own clock: say so on the session
                lines = [cal.time_label(d, p, sessions[owner]["length"]) or "", *lines]
            cell.value = "\n".join(x for x in lines if x)
            if owner is not None:
                cell.fill = PatternFill("solid", fgColor=course_hex(sessions[owner]["course_code"]))
        if p in cal.break_after:
            for c in range(1, ncols + 1):
                cur = ws.cell(row=row, column=c).border
                ws.cell(row=row, column=c).border = Border(
                    left=cur.left, right=cur.right, top=cur.top, bottom=Side(style="medium", color="444444"))

    for i in merge:
        s = sessions[i]
        r0 = s["start"] + 3
        ws.merge_cells(start_row=r0, start_column=s["day"] + 2, end_row=r0 + s["length"] - 1,
                       end_column=s["day"] + 2)

    ws.column_dimensions["A"].width = 6
    for d in range(cal.days):
        ws.column_dimensions[get_column_letter(d + 2)].width = 28
    ws.freeze_panes = "B3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def build_xlsx(inst: Institution, tt: Timetable, out: BinaryIO) -> None:
    cal = inst.calendar
    wb = Workbook()
    flat = wb.active
    flat.title = "All sessions"
    flat.append(["Day", "Lecture", "Course code", "Course", "Faculty", "Batches", "Room", "Offering"])
    for c in flat[1]:
        c.fill, c.font = _HEAD_FILL, _HEAD_FONT
    rows = []
    for eid, _ in entities(inst, "room"):
        rows += session_views(inst, tt, "room", eid)  # every session appears in exactly one room
    for s in sorted(rows, key=lambda s: (s["day"], s["start"], s["course_code"])):
        flat.append([s["day_name"], lecture_label(s), s["course_code"], s["course_name"],
                     s["faculty_name"], ", ".join(s["batch_ids"]), s["room_name"], s["offering_id"]])
    for col, w in zip("ABCDEFGH", (8, 9, 16, 34, 24, 26, 20, 18)):
        flat.column_dimensions[col].width = w
    flat.freeze_panes = "A2"

    taken = {"all sessions"}
    for kind, eid, label, sessions in _selection(inst, tt, None, None):
        ws = wb.create_sheet(_sheet_name(KIND_TITLES[kind], eid, taken))
        _write_grid(ws, cal, kind, f"{inst.name} · {KIND_TITLES[kind]}: {label}", sessions)
    wb.save(out)


# ---------- PDF ----------

def _register_font(font_path: str) -> str:
    """Register a TrueType font (.ttf, or the first face of a .ttc) and return its name."""
    name = "AAFont"
    pdfmetrics.registerFont(TTFont(name, font_path, subfontIndex=0))
    pdfmetrics.registerFontFamily(name, normal=name, bold=name, italic=name, boldItalic=name)
    return name


def _beyond_latin1(inst: Institution) -> bool:
    return any(ord(ch) > 255 for ch in inst.model_dump_json(ensure_ascii=False))


def build_pdf(inst: Institution, tt: Timetable, out: BinaryIO,
              kind: str | None = None, ident: str | None = None, font_path: str | None = None) -> None:
    """Landscape A4, one page per class / faculty member / room.

    The built-in PDF fonts only draw Latin text. For Hindi or other Indian scripts pass
    `font_path` to a TrueType font that has them (the API reads it from AA_PDF_FONT).
    """
    cal = inst.calendar
    sel = _selection(inst, tt, kind, ident)

    base_font = "Helvetica"
    if font_path:
        base_font = _register_font(font_path)
    elif _beyond_latin1(inst):
        log.warning("institution data contains non-Latin text but no PDF font is configured; "
                    "it will print as boxes. Set AA_PDF_FONT to a TrueType font that covers it.")

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Heading2"], spaceAfter=2,
                                 fontName=base_font if font_path else "Helvetica-Bold")
    sub_style = ParagraphStyle("s", parent=styles["Normal"], fontName=base_font, fontSize=8,
                               textColor=colors.grey, spaceAfter=6)
    cell = ParagraphStyle("c", parent=styles["Normal"], fontName=base_font, fontSize=7, leading=8.4)
    tiny = ParagraphStyle("tiny", parent=cell, fontSize=6, leading=7)
    head = ParagraphStyle("h", parent=cell, fontSize=8, alignment=1, textColor=colors.white)

    page_w, page_h = landscape(A4)
    margin = 1 * cm
    usable_w = page_w - 2 * margin
    row_h = min(2.3 * cm, (page_h - 2 * margin - 3.2 * cm) / cal.lectures_per_day)
    time_w = 1.2 * cm
    col_w = (usable_w - time_w) / max(cal.days, 1)

    story: list = []
    if not sel:
        story.append(Paragraph("No sessions to export.", styles["Normal"]))

    for n, (k, eid, label, sessions) in enumerate(sel):
        if n:
            story.append(PageBreak())
        story.append(Paragraph(escape(f"{KIND_TITLES[k]}: {label}"), title_style))
        story.append(Paragraph(escape(inst.name), sub_style))

        cover, merge = layout(sessions)
        merged = set(merge)
        starts = {(s["day"], s["start"]): i for i, s in enumerate(sessions)}

        data: list[list] = [[""] + [Paragraph(escape(d), head) for d in cal.day_names]]
        style = [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("VALIGN", (0, 1), (0, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]
        for p in range(cal.lectures_per_day):
            r = p + 1
            shown = cal.times[p] if cal.times else None
            when = (f"<br/>{LectureTime.pretty(shown.start)}\u2013{LectureTime.pretty(shown.end)}" if shown else "")
            row: list = [Paragraph(f"<b>L{p + 1}</b>{when}", ParagraphStyle("p", parent=cell, alignment=1))]
            for d in range(cal.days):
                idxs = cover.get((d, p), [])
                owner = idxs[0] if len(idxs) == 1 else None
                first = owner is not None and starts.get((d, p)) == owner
                if not idxs or (owner in merged and not first):
                    row.append("")
                    continue
                lines = _cell_text(k, sessions, idxs, first)  # type: ignore[arg-type]
                if first and d in cal.day_times:
                    lines = [cal.time_label(d, p, sessions[owner]["length"]) or "", *lines]
                    lines = [x for x in lines if x]
                row.append(Paragraph("<br/>".join(escape(x) for x in lines), cell if len(idxs) == 1 else tiny))
                if owner is not None:
                    rgb = course_rgb(sessions[owner]["course_code"])
                    style.append(("BACKGROUND", (d + 1, r), (d + 1, r), colors.Color(*rgb)))
            data.append(row)
            if p in cal.break_after:
                style.append(("LINEBELOW", (0, r), (-1, r), 1.6, colors.HexColor("#444444")))
        for i in merge:
            s = sessions[i]
            c, r0 = s["day"] + 1, s["start"] + 1
            style.append(("SPAN", (c, r0), (c, r0 + s["length"] - 1)))
            style.append(("BACKGROUND", (c, r0), (c, r0 + s["length"] - 1),
                          colors.Color(*course_rgb(s["course_code"]))))

        tbl = Table(data, colWidths=[time_w] + [col_w] * cal.days,
                    rowHeights=[0.6 * cm] + [row_h] * cal.lectures_per_day)
        tbl.setStyle(TableStyle(style))
        story.append(tbl)
        story.append(Spacer(1, 4))

    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=margin, title=f"{inst.name} timetable",
                            author="Automated Academics")
    doc.build(story)
