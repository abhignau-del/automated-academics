"""Import a *workload list* into an Institution.

Most colleges already keep one flat table: for every subject, who teaches it, to which programme
and semester, how many hours a week, and how many students. That is nearly all the solver needs, so
this reads it directly instead of asking people to rebuild it as the six-sheet template.

Columns are found by heading (any order, case and punctuation ignored):

  subject   Subject, Paper, Course name, Title
  faculty   Faculty, Faculty name, Teacher, Instructor
  program   Programme, Program, Course, Degree, Class
  sem       Semester, Sem
  division  Division, Div, Section                         (optional, default 1)
  load      Load per week, Hours per week, Lectures per week, Actual load
  students  Total students, Students, Strength, Intake      (optional)
  credit    Credit, Credits                                 (optional)

Whatever the sheet does not say (rooms, the week's shape, who is a lab course) is filled with
stated assumptions, returned in the report so nothing is silent.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook

from .models import Batch, Calendar, Course, Faculty, Institution, Level, Offering, Room, RoomKind

# heading -> aliases in priority order (the first alias that is present wins)
ALIASES: dict[str, list[str]] = {
    "subject": ["subject", "subjectname", "paper", "papername", "coursename", "title"],
    "faculty": ["faculty", "facultyname", "teacher", "teachername", "instructor"],
    "program": ["program", "programme", "course", "degree", "class", "programname"],
    "sem": ["sem", "semester"],
    "division": ["division", "div", "section"],
    "load": ["loadperweek", "hoursperweek", "lecturesperweek", "periodsperweek", "load", "actualload", "hours"],
    "students": ["totalstudents", "students", "strength", "intake", "noofstudents"],
    "credit": ["credit", "credits"],
}
REQUIRED = ("subject", "faculty", "program", "load")


@dataclass
class WorkloadOptions:
    days: list[str] = field(default_factory=lambda: ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"])
    lectures_per_day: int = 4
    break_after: list[int] = field(default_factory=lambda: [2])  # 1-based: a break after the 2nd lecture
    classrooms: int | None = None  # default: one per class
    seats: int | None = None  # default: the biggest group, rounded up to 10
    labs: int = 0
    lab_seats: int | None = None
    # Teach the same subject jointly when the same teacher has it in several programmes of one
    # semester and the combined class stays within this many students. 0 switches it off.
    joint_max_students: int = 0
    default_students: int = 60  # for rows with no student count


@dataclass
class WorkloadReport:
    rows_read: int = 0
    rows_used: int = 0
    assumptions: list[str] = field(default_factory=list)
    joint: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)  # things in the sheet worth a second look

    def as_dict(self) -> dict[str, Any]:
        return {"rows_read": self.rows_read, "rows_used": self.rows_used, "assumptions": self.assumptions,
                "joint": self.joint, "problems": self.problems}


class WorkloadError(ValueError):
    """The sheet could not be read as a workload list."""


def _norm(s: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def _clean(s: Any) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def _num(v: Any) -> int | None:
    try:
        f = float(str(v).strip())
    except (TypeError, ValueError):
        return None
    return int(f) if f == int(f) else None


def _initials(name: str, taken: set[str]) -> str:
    parts = [p for p in re.split(r"[\s.]+", name) if p and p.lower() not in ("dr", "prof", "mr", "mrs", "ms")]
    base = ("".join(p[0] for p in parts[:2]) or "F").upper()
    code, n = base, 2
    while code in taken:
        code, n = f"{base}{n}", n + 1
    taken.add(code)
    return code


def _code_for(subject: str, taken: set[str]) -> str:
    words = [w for w in re.split(r"[^A-Za-z0-9]+", subject) if w and w.lower() not in ("of", "and", "the", "to", "in", "&")]
    base = ("".join(w[0] for w in words[:5]).upper() or "C")
    if len(base) < 2:
        base = (re.sub(r"[^A-Za-z0-9]", "", subject).upper()[:4] or "C")
    code, n = base, 2
    while code in taken:
        code, n = f"{base}{n}", n + 1
    taken.add(code)
    return code


def _read_rows(source: str | Path | BinaryIO, sheet: str | None) -> tuple[list[dict[str, Any]], WorkloadReport]:
    report = WorkloadReport()
    wb = load_workbook(source, data_only=True, read_only=True)
    names = [sheet] if sheet else wb.sheetnames
    best: tuple[int, Any, int, dict[str, int]] | None = None  # (score, worksheet, header row, columns)
    for name in names:
        if name not in wb.sheetnames:
            raise WorkloadError(f"sheet {name!r} not found")
        ws = wb[name]
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
            heads = {_norm(v): i for i, v in enumerate(row) if v is not None}
            cols: dict[str, int] = {}
            for key, aliases in ALIASES.items():
                hit = next((a for a in aliases if a in heads), None)
                if hit is not None:
                    cols[key] = heads[hit]
            if all(k in cols for k in REQUIRED):
                score = len(cols)
                if best is None or score > best[0]:
                    best = (score, ws, r, cols)
    if best is None:
        raise WorkloadError("could not find a heading row with Subject, Faculty, Programme and Load columns")
    _, ws, header, cols = best
    rows: list[dict[str, Any]] = []
    for r, row in enumerate(ws.iter_rows(min_row=header + 1, values_only=True), start=header + 1):
        if all(v in (None, "") for v in row):
            continue
        rec = {k: (row[i] if i < len(row) else None) for k, i in cols.items()}
        rec["_row"] = r
        if not any(_clean(rec.get(k)) for k in ("subject", "faculty", "program")):
            continue  # a numbered placeholder row with nothing on it: not worth mentioning
        report.rows_read += 1
        rows.append(rec)
    report.assumptions.append(f"Read sheet \"{ws.title}\", heading on row {header}.")
    return rows, report


def import_workload(source: str | Path | BinaryIO, options: WorkloadOptions | None = None,
                    sheet: str | None = None) -> tuple[Institution, WorkloadReport]:
    """Build an institution from a workload list. Raises WorkloadError if it can't be read at all."""
    opt = options or WorkloadOptions()
    rows, report = _read_rows(source, sheet)

    clean: list[dict[str, Any]] = []
    for rec in rows:
        subject, fac, prog = _clean(rec["subject"]), _clean(rec["faculty"]), _clean(rec["program"])
        load = _num(rec["load"])
        if not subject or not fac or not prog:
            report.problems.append(f"row {rec['_row']}: skipped, it has no {'subject' if not subject else 'faculty' if not fac else 'programme'}")
            continue
        if not load or load < 1:
            report.problems.append(f"row {rec['_row']}: skipped, \"{subject}\" ({prog}) has no hours per week")
            continue
        clean.append({**rec, "subject": subject, "faculty": fac, "program": prog, "load": load})
    report.rows_used = len(clean)
    if not clean:
        raise WorkloadError("no usable rows: every row was missing a subject, teacher, programme or hours")

    # ---- classes: one per (programme, semester, division)
    batch_key: dict[tuple, str] = {}
    batches: list[Batch] = []
    strengths: dict[str, int] = {}
    divs: dict[tuple, set[str]] = defaultdict(set)
    for r in clean:
        divs[(r["program"], _num(r.get("sem")) or 1)].add(_clean(r.get("division") or 1))
    several = {k for k, v in divs.items() if len(v) > 1}  # only these need the division in their id
    for r in clean:
        sem, div = _num(r.get("sem")) or 1, _clean(r.get("division") or 1)
        key = (r["program"], sem, div)
        if key in batch_key:
            continue
        bid = f"{re.sub(r'[^A-Za-z0-9]+', '', r['program'])}-S{sem}" + (f"-{div}" if (r["program"], sem) in several else "")
        n = _num(r.get("students"))
        if not n:
            n = opt.default_students
            report.problems.append(f"{bid}: no student count in the sheet, assumed {n}")
        batch_key[key] = bid
        strengths[bid] = n
        level = Level.PG if re.match(r"^(m|post|pg)", r["program"], re.I) else Level.UG
        batches.append(Batch(id=bid, program=r["program"], level=level, semester=sem, section=div,
                             department=r["program"], strength=n))
    report.assumptions.append("Programmes starting with M (MBA, MCA, MSc...) are treated as PG, the rest as UG; "
                              "each programme is also used as its department. Edit these in the Data tab if needed.")

    # ---- courses and faculty
    course_codes: dict[str, str] = {}
    courses: list[Course] = []
    taken_codes: set[str] = set()
    for r in clean:
        k = _norm(r["subject"])
        if k not in course_codes:
            code = _code_for(r["subject"], taken_codes)
            course_codes[k] = code
            courses.append(Course(code=code, name=r["subject"], department=r["program"], credits=_num(r.get("credit")) or 0))
    faculty_ids: dict[str, str] = {}
    faculty: list[Faculty] = []
    taken_ids: set[str] = set()
    for r in clean:
        k = _norm(r["faculty"])
        if k not in faculty_ids:
            fid = _initials(r["faculty"], taken_ids)
            faculty_ids[k] = fid
            faculty.append(Faculty(id=fid, name=r["faculty"], department=r["program"]))
    report.assumptions.append("Every teacher is assumed to be available at all times (the sheet has no availability).")

    # ---- offerings, optionally joining classes that share a teacher and subject
    groups: dict[tuple, list[dict[str, Any]]] = defaultdict(list)
    for r in clean:
        groups[(_norm(r["subject"]), _norm(r["faculty"]), _num(r.get("sem")) or 1)].append(r)

    offerings: list[Offering] = []
    used_ids: set[str] = set()

    def add_offering(rows_: list[dict[str, Any]]) -> None:
        code = course_codes[_norm(rows_[0]["subject"])]
        loads = {r["load"] for r in rows_}
        load = max(loads)
        bids = [batch_key[(r["program"], _num(r.get("sem")) or 1, _clean(r.get("division") or 1))] for r in rows_]
        if len(loads) > 1:
            detail = ", ".join(f"{b} {r['load']}" for b, r in zip(bids, rows_))
            report.problems.append(f"{rows_[0]['subject']}: hours differ between classes taught together ({detail}); used {load}")
        if len(rows_) > 1:
            report.joint.append(f"{rows_[0]['subject']} ({faculty_ids[_norm(rows_[0]['faculty'])]}): {' + '.join(bids)} "
                                f"= {sum(strengths[b] for b in bids)} students")
        oid = f"O-{code}-" + "+".join(bids) if len(bids) <= 2 else f"O-{code}-{bids[0]}+{len(bids) - 1}"
        n = 2
        base = oid
        while oid in used_ids:
            oid, n = f"{base}~{n}", n + 1
        used_ids.add(oid)
        offerings.append(Offering(id=oid, course_code=code, faculty_id=faculty_ids[_norm(rows_[0]["faculty"])],
                                  batch_ids=bids, sessions=[1] * load))

    for rows_ in groups.values():
        if opt.joint_max_students > 0 and len({r["program"] for r in rows_}) > 1:
            current: list[dict[str, Any]] = []
            size = 0
            for r in rows_:
                s = strengths[batch_key[(r["program"], _num(r.get("sem")) or 1, _clean(r.get("division") or 1))]]
                if current and size + s > opt.joint_max_students:
                    add_offering(current)
                    current, size = [], 0
                current.append(r)
                size += s
            if current:
                add_offering(current)
        else:
            for r in rows_:
                add_offering([r])
    if opt.joint_max_students > 0:
        report.assumptions.append(f"Classes with the same teacher and subject are combined while they total "
                                  f"{opt.joint_max_students} students or fewer ({len(report.joint)} combined groups).")
    else:
        report.assumptions.append("Each programme's row is its own offering (no joint teaching).")
    report.assumptions.append("Every session is one lecture long. The sheet does not say which subjects are labs; "
                              "set Needs = lab on practical courses in the Data tab.")

    # ---- rooms the sheet doesn't describe
    biggest = max(sum(strengths[b] for b in o.batch_ids) for o in offerings)  # the largest class that sits together
    seats = opt.seats or -(-biggest // 10) * 10
    n_rooms = opt.classrooms if opt.classrooms is not None else max(2, len(batches))
    rooms = [Room(id=f"R{i}", name=f"Classroom {i}", capacity=seats) for i in range(1, n_rooms + 1)]
    rooms += [Room(id=f"L{i}", name=f"Lab {i}", capacity=opt.lab_seats or seats, kind=RoomKind.LAB) for i in range(1, opt.labs + 1)]
    report.assumptions.append(f"The sheet has no rooms, so {n_rooms} classroom(s) of {seats} seats"
                              + (f" and {opt.labs} lab(s)" if opt.labs else "") + " were assumed.")

    cal = Calendar(day_names=opt.days, lectures_per_day=opt.lectures_per_day,
                   break_after=[b - 1 for b in opt.break_after if 1 <= b < opt.lectures_per_day])
    report.assumptions.append(f"The week is {len(opt.days)} days of {opt.lectures_per_day} lectures"
                              + (f" with a break after lecture {', '.join(map(str, opt.break_after))}." if opt.break_after else "."))
    inst = Institution(name="Imported workload", calendar=cal, rooms=rooms, faculty=faculty, batches=batches,
                       courses=courses, offerings=offerings)
    return inst, report
