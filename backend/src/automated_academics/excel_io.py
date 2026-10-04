"""Excel import/export of an Institution.

Workbook layout (one header row, then data rows; header names are case-insensitive):

  Institution : key | value        (name, day_names, lectures_per_day, break_after)
  Rooms       : id, name, capacity, kind
  Faculty     : id, name, department, max_lectures_per_day, unavailable
  Batches     : id, program, level, semester, section, department, strength, group_of
  Courses     : code, name, department, credits, category, room_kind
  Offerings   : id, course_code, faculty_id, batch_ids, sessions

For humans, lectures are numbered from 1 and days are written by name:
  unavailable = "Mon:1, Tue:3"   batch_ids = "CS-UG1, MG-UG1"   sessions = "1,1,1" or "2"
`break_after` is likewise 1-based: "4" means a break after the 4th lecture.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, BinaryIO, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from .models import (
    Batch,
    Calendar,
    Course,
    Faculty,
    Institution,
    Level,
    Offering,
    Room,
    RoomKind,
    Slot,
)

SHEETS: dict[str, list[str]] = {
    "Rooms": ["id", "name", "capacity", "kind"],
    "Faculty": ["id", "name", "department", "max_lectures_per_day", "unavailable"],
    "Batches": ["id", "program", "level", "semester", "section", "department", "strength", "group_of"],
    "Courses": ["code", "name", "department", "credits", "category", "room_kind"],
    "Offerings": ["id", "course_code", "faculty_id", "batch_ids", "sessions"],
}
REQUIRED: dict[str, set[str]] = {
    "Rooms": {"id", "name", "capacity"},
    "Faculty": {"id", "name", "department"},
    "Batches": {"id", "program", "level", "semester", "department", "strength"},
    "Courses": {"code", "name", "department", "credits"},
    "Offerings": {"id", "course_code", "faculty_id", "batch_ids", "sessions"},
}


@dataclass(frozen=True)
class ImportIssue:
    sheet: str
    row: int | None  # 1-based Excel row, None for sheet-level problems
    message: str

    def __str__(self) -> str:
        where = self.sheet if self.row is None else f"{self.sheet} row {self.row}"
        return f"{where}: {self.message}"


class ImportErrors(ValueError):
    """Raised with every problem found, so users can fix the sheet in one pass."""

    def __init__(self, issues: list[ImportIssue]):
        self.issues = issues
        super().__init__("\n".join(str(i) for i in issues))


# ---------- parsing helpers ----------

def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _int(v: Any, field: str) -> int:
    s = _text(v)
    try:
        f = float(s)
    except ValueError:
        raise ValueError(f"{field} must be a whole number, got {s!r}") from None
    if f != int(f):
        raise ValueError(f"{field} must be a whole number, got {s!r}")
    return int(f)


def _split(v: Any) -> list[str]:
    return [p.strip() for p in _text(v).replace(";", ",").split(",") if p.strip()]


def _enum(cls: type, v: Any, field: str, default: Any = None):
    s = _text(v)
    if not s:
        if default is not None:
            return default
        raise ValueError(f"{field} is required")
    for member in cls:
        if s.lower() == member.value.lower():
            return member
    allowed = ", ".join(m.value for m in cls)
    raise ValueError(f"{field} must be one of [{allowed}], got {s!r}")


# ---------- reading ----------

def _rows(ws, sheet: str, issues: list[ImportIssue]):
    """Yield (excel_row, {column: raw_value}) for non-blank data rows."""
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        issues.append(ImportIssue(sheet, None, "sheet is empty"))
        return
    cols = [_text(h).lower() for h in header]
    missing = REQUIRED[sheet] - set(cols)
    if missing:
        issues.append(ImportIssue(sheet, 1, f"missing column(s): {', '.join(sorted(missing))}"))
        return
    for n, row in enumerate(it, start=2):
        if all(_text(c) == "" for c in row):
            continue
        yield n, dict(zip(cols, row))


def _collect(ws, sheet: str, issues: list[ImportIssue], build: Callable[[dict], Any]) -> list:
    out = []
    for n, rec in _rows(ws, sheet, issues):
        try:
            out.append(build(rec))
        except (ValueError, ValidationError) as e:
            msg = "; ".join(err["msg"] for err in e.errors()) if isinstance(e, ValidationError) else str(e)
            issues.append(ImportIssue(sheet, n, msg))
    return out


def _calendar(ws, issues: list[ImportIssue]) -> tuple[str, Calendar]:
    kv = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and _text(row[0]):
            kv[_text(row[0]).lower()] = row[1] if len(row) > 1 else None
    name = _text(kv.get("name")) or "Institution"
    try:
        days = _split(kv.get("day_names")) or Calendar().day_names
        lectures = _int(kv.get("lectures_per_day") or 7, "lectures_per_day")
        breaks = [_int(b, "break_after") - 1 for b in _split(kv.get("break_after"))]
        return name, Calendar(day_names=days, lectures_per_day=lectures, break_after=breaks)
    except (ValueError, ValidationError) as e:
        issues.append(ImportIssue("Institution", None, str(e)))
        return name, Calendar()


def import_workbook(path: str | Path | BinaryIO) -> Institution:
    wb = load_workbook(path, data_only=True)
    issues: list[ImportIssue] = []

    absent = [s for s in ["Institution", *SHEETS] if s not in wb.sheetnames]
    if absent:
        raise ImportErrors([ImportIssue(s, None, "sheet not found") for s in absent])

    name, cal = _calendar(wb["Institution"], issues)
    day_index = {d.lower(): i for i, d in enumerate(cal.day_names)}

    def slots(v: Any) -> list[Slot]:
        out = []
        for tok in _split(v):
            day, _, per = tok.partition(":")
            if day.strip().lower() not in day_index or not per.strip():
                raise ValueError(f"unavailable entry {tok!r} must look like 'Mon:1'")
            p = _int(per, "unavailable lecture")
            if not 1 <= p <= cal.lectures_per_day:
                raise ValueError(f"unavailable lecture {p} outside 1..{cal.lectures_per_day}")
            out.append(Slot(day=day_index[day.strip().lower()], lecture=p - 1))
        return out

    rooms = _collect(wb["Rooms"], "Rooms", issues, lambda r: Room(
        id=_text(r["id"]), name=_text(r["name"]), capacity=_int(r["capacity"], "capacity"),
        kind=_enum(RoomKind, r.get("kind"), "kind", RoomKind.CLASSROOM)))

    faculty = _collect(wb["Faculty"], "Faculty", issues, lambda r: Faculty(
        id=_text(r["id"]), name=_text(r["name"]), department=_text(r["department"]),
        max_lectures_per_day=_int(r.get("max_lectures_per_day") or 6, "max_lectures_per_day"),
        unavailable=slots(r.get("unavailable"))))

    batches = _collect(wb["Batches"], "Batches", issues, lambda r: Batch(
        id=_text(r["id"]), program=_text(r["program"]), level=_enum(Level, r["level"], "level"),
        semester=_int(r["semester"], "semester"), section=_text(r.get("section")) or "A",
        department=_text(r["department"]), strength=_int(r["strength"], "strength"),
        group_of=_text(r.get("group_of")) or None))

    courses = _collect(wb["Courses"], "Courses", issues, lambda r: Course(
        code=_text(r["code"]), name=_text(r["name"]), department=_text(r["department"]),
        credits=_int(r["credits"], "credits"), category=_text(r.get("category")) or "DSC",
        room_kind=_enum(RoomKind, r.get("room_kind"), "room_kind", RoomKind.CLASSROOM)))

    offerings = _collect(wb["Offerings"], "Offerings", issues, lambda r: Offering(
        id=_text(r["id"]), course_code=_text(r["course_code"]), faculty_id=_text(r["faculty_id"]),
        batch_ids=_split(r["batch_ids"]),
        sessions=[_int(s, "sessions") for s in _split(r["sessions"])]))

    for sheet, items, key in [("Rooms", rooms, "id"), ("Faculty", faculty, "id"),
                              ("Batches", batches, "id"), ("Courses", courses, "code"),
                              ("Offerings", offerings, "id")]:
        seen: set[str] = set()
        for it in items:
            k = getattr(it, key)
            if k in seen:
                issues.append(ImportIssue(sheet, None, f"duplicate {key} {k!r}"))
            seen.add(k)

    if issues:
        raise ImportErrors(issues)
    try:
        return Institution(name=name, calendar=cal, rooms=rooms, faculty=faculty,
                           batches=batches, courses=courses, offerings=offerings)
    except ValidationError as e:  # cross-sheet references
        raise ImportErrors([ImportIssue("Offerings/Batches", None, err["msg"].removeprefix("Value error, "))
                            for err in e.errors()]) from None


# ---------- writing (templates and exports) ----------

_HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
_HEAD_FONT = Font(bold=True, color="FFFFFF")


def _style(ws, widths: dict[int, int] | None = None) -> None:
    for c in ws[1]:
        c.fill, c.font = _HEAD_FILL, _HEAD_FONT
        c.alignment = Alignment(vertical="center")
    for i, col in enumerate(ws.iter_cols(), start=1):
        longest = max((len(_text(c.value)) for c in col), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, longest + 2), 48)
    ws.freeze_panes = "A2"


def _dropdown(ws, header: str, options: list[str]) -> None:
    col = next(i for i, c in enumerate(ws[1], start=1) if c.value == header)
    letter = get_column_letter(col)
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}1000")


def export_workbook(inst: Institution, path: str | Path | BinaryIO) -> None:
    """Write an Institution as a workbook the importer can read back."""
    cal = inst.calendar
    wb = Workbook()
    ws = wb.active
    ws.title = "Institution"
    ws.append(["key", "value"])
    ws.append(["name", inst.name])
    ws.append(["day_names", ", ".join(cal.day_names)])
    ws.append(["lectures_per_day", cal.lectures_per_day])
    ws.append(["break_after", ", ".join(str(b + 1) for b in cal.break_after)])
    _style(ws)

    def sheet(title: str, rows: list[list]) -> Any:
        w = wb.create_sheet(title)
        w.append(SHEETS[title])
        for r in rows:
            w.append(r)
        _style(w)
        return w

    sheet("Rooms", [[r.id, r.name, r.capacity, r.kind.value] for r in inst.rooms])
    sheet("Faculty", [[f.id, f.name, f.department, f.max_lectures_per_day,
                       ", ".join(f"{cal.day_names[s.day]}:{s.lecture + 1}" for s in f.unavailable)]
                      for f in inst.faculty])
    sheet("Batches", [[b.id, b.program, b.level.value, b.semester, b.section, b.department,
                       b.strength, b.group_of or ""] for b in inst.batches])
    sheet("Courses", [[c.code, c.name, c.department, c.credits, c.category, c.room_kind.value]
                      for c in inst.courses])
    sheet("Offerings", [[o.id, o.course_code, o.faculty_id, ", ".join(o.batch_ids),
                         ", ".join(str(s) for s in o.sessions)] for o in inst.offerings])

    _dropdown(wb["Rooms"], "kind", [k.value for k in RoomKind])
    _dropdown(wb["Courses"], "room_kind", [k.value for k in RoomKind])
    _dropdown(wb["Batches"], "level", [lv.value for lv in Level])
    wb.save(path)
