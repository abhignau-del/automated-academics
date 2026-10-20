"""Fictional timetable grids for the browser journeys.

    python make_grid_fixture.py ordinary.xlsx          two classes (one with parallel electives), a repeated week on a second sheet
    python make_grid_fixture.py sideways.xlsx --sideways   one class drawn with the days down the side
"""

import sys

from openpyxl import Workbook

DAYS = ["MON", "TUE", "WED", "THU", "FRI", "SAT"]
TIMES = ["9 TO 9.50", "9.55 TO 10.45", "10.50 TO 11.40", "12 TO 12.50"]

ECON1 = [
    ["Micro(Prof. Ira Bose)", "Micro(Prof. Ira Bose)", "Statistics(Dr. Omar Lal)", "", "Macro(Prof. Tia Sen)", ""],
    ["Macro(Prof. Tia Sen)", "Statistics(Dr. Omar Lal)", "", "Micro(Prof. Ira Bose)", "Statistics(Dr. Omar Lal)", "Macro(Prof. Tia Sen)"],
    ["", "Macro(Prof. Tia Sen)", "Micro(Prof. Ira Bose)", "Statistics(Dr. Omar Lal)", "", "Micro(Prof. Ira Bose)"],
    ["Statistics(Dr. Omar Lal)", "", "Macro(Prof. Tia Sen)", "Macro(Prof. Tia Sen)", "Micro(Prof. Ira Bose)", ""],
]  # lectures x days
ECON3 = [
    ["Micro(Prof. Ira Bose)", "Econometrics(Dr. Omar Lal)", "", "Econometrics(Dr. Omar Lal)", "", "Policy(Prof. Tia Sen)"],
    ["Policy(Prof. Tia Sen)", "", "Econometrics(Dr. Omar Lal)", "", "Policy(Prof. Tia Sen)", ""],
    ["1) Trade (Prof. Ira Bose) 2) Finance (Dr. Omar Lal)", "", "", "1) Trade (Prof. Ira Bose) 2) Finance (Dr. Omar Lal)", "", "Econometrics(Dr. Omar Lal)"],
    ["", "Policy(Prof. Tia Sen)", "", "", "", ""],
]


def put(ws, top, title, rows, times=TIMES):
    ws.cell(row=top, column=1, value=title)
    ws.cell(row=top + 1, column=3, value="FROM 4TH MARCH 2030")
    for j, d in enumerate(DAYS):
        ws.cell(row=top + 2, column=2 + j, value=d)
    for i, row in enumerate(rows):
        ws.cell(row=top + 3 + i, column=1, value=times[i])
        for j, v in enumerate(row):
            if v:
                ws.cell(row=top + 3 + i, column=2 + j, value=v)
    return top + 4 + len(rows)


def put_sideways(ws, top, title, rows, times=TIMES):
    ws.cell(row=top, column=1, value=title)
    for j, tm in enumerate(times):
        ws.cell(row=top, column=2 + j, value=tm)
    for d in range(len(DAYS)):
        ws.cell(row=top + 1 + d, column=1, value=DAYS[d])
        for lecture, row in enumerate(rows):
            if row[d]:
                ws.cell(row=top + 1 + d, column=2 + lecture, value=row[d])


wb = Workbook()
wb.remove(wb.active)
if "--sideways" in sys.argv:
    put_sideways(wb.create_sheet("Term 1"), 1, "BA Economics Sem 1 ( MAIN BLOCK, ROOM NO_21)", ECON1)
else:
    ws = wb.create_sheet("Term 1")
    nxt = put(ws, 1, "BA Economics Sem 1 ( MAIN BLOCK, ROOM NO_21)", ECON1)
    put(ws, nxt + 1, "BA Economics Sem 3 ( MAIN BLOCK, ROOM NO_22)", ECON3)
    put(wb.create_sheet("Term 1 (short week)"), 1, "BA Economics Sem 1", ECON1,
        times=["9 TO 9.40", "9.45 TO 10.25", "10.30 TO 11.10", "11.30 TO 12.10"])
wb.save(sys.argv[1])
print("saved", sys.argv[1])
